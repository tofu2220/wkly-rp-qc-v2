import hashlib
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from wkly_rp_qc_v2.importer import inspect_workbook, normalized_payload


def write_workbook(path: Path, template: str, *, draft: bool = False) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BC Tuần 01.09 - 05.09"
    if template == "early":
        for cell, value in {
            "A4": "STT", "B4": "Sản phẩm", "C4": "Kết quả tuần 01.09 - 05.09",
            "D4": "Kế hoạch tuần 07.09 - 12.09", "E4": "Vấn đề/ rủi ro",
            "F4": "PIC", "A5": "1", "B5": "Project Early", "A6": "1.1",
            "B6": "PTSP", "F6": "Owner", "C6": "Actual line 1\nActual line 2",
            "D6": "Deliver feature", "E6": "Blocker text",
        }.items():
            sheet[cell] = value
    elif template == "revised":
        for cell, value in {
            "B5": "STT", "C5": "VAI TRÒ", "D5": "Nhóm vai trò", "E5": "KẾT QUẢ TUẦN",
            "F5": "KẾ HOẠCH TUẦN", "G5": "VẤN ĐỀ / RỦI RO", "H5": "PIC",
            "B6": "II", "C6": "Project Revised", "B7": "2.1", "D7": "PTSP",
            "E7": "Actual revised", "F7": "Plan revised", "G7": "Risk revised", "H7": "Owner",
        }.items():
            sheet[cell] = value
        sheet.merge_cells("C6:C7")
    elif template == "latest":
        for cell, value in {
            "B5": "Dự án", "C5": "STT", "D5": "VAI TRÒ", "F5": "KẾT QUẢ TUẦN",
            "G5": "KẾ HOẠCH TUẦN", "H5": "VẤN ĐỀ / RỦI RO", "I5": "PIC",
            "B6": "II", "D6": "Project Latest", "C7": "2.1", "D7": "PTSP",
            "F7": "Actual latest", "G7": "Plan latest", "H7": "Risk latest", "I7": "Owner",
        }.items():
            sheet[cell] = value
    elif template == "unknown":
        for cell, value in {
            "A1": "Unrecognized layout", "A2": "Different heading", "B7": "Unknown project",
        }.items():
            sheet[cell] = value
    else:
        raise ValueError(template)

    if draft and template != "unknown":
        narrative_cells = {
            "early": ("C6", "D6", "E6"),
            "revised": ("E7", "F7", "G7"),
            "latest": ("F7", "G7", "H7"),
        }[template]
        for cell in narrative_cells:
            sheet[cell] = None
    workbook.save(path)
    workbook.close()


class WorkbookImporterTests(unittest.TestCase):
    def test_inspect_workbook_recognizes_early_revised_and_latest_layouts(self):
        for template in ("early", "revised", "latest"):
            with self.subTest(template=template), tempfile.TemporaryDirectory() as temp:
                path = Path(temp) / f"{template}.xlsx"
                write_workbook(path, template)

                preview = inspect_workbook(path)

                self.assertEqual(preview.workbook_sha256, hashlib.sha256(path.read_bytes()).hexdigest())
                self.assertEqual(preview.sheets[0]["template"], template)
                self.assertEqual(preview.sheets[0]["mapping_status"], "recognized")
                self.assertEqual(len(preview.rows), 1)
                self.assertEqual(preview.validation["status"], "PASS")

    def test_inspect_workbook_preserves_merged_and_hierarchical_provenance(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "revised.xlsx"
            write_workbook(path, "revised")

            preview = inspect_workbook(path)
            row = preview.rows[0]

            self.assertEqual(row["project"], "Project Revised")
            self.assertEqual(row["role"], "PTSP")
            self.assertEqual(row["source_refs"]["project"], "C6")
            self.assertEqual(row["source_refs"]["role"], "D7")
            self.assertEqual(row["field_provenance"]["project"]["merged_range"], "C6:C7")
            self.assertEqual(row["field_provenance"]["project"]["resolution"], "hierarchical_project_group")

    def test_inspect_workbook_keeps_draft_rows_and_marks_draft_sheet(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "draft.xlsx"
            write_workbook(path, "latest", draft=True)

            preview = inspect_workbook(path)

            self.assertEqual(len(preview.rows), 1)
            self.assertEqual(preview.rows[0]["actual"], None)
            self.assertEqual(preview.rows[0]["next_plan"], None)
            self.assertEqual(preview.sheets[0]["sheet_state"], "draft_or_incomplete")
            self.assertEqual(preview.validation["status"], "PASS")

    def test_inspect_workbook_returns_mapping_required_for_unknown_layout(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "unknown.xlsx"
            write_workbook(path, "unknown")

            preview = inspect_workbook(path)

            self.assertEqual(preview.sheets[0]["mapping_status"], "mapping_required")
            self.assertIsNone(preview.sheets[0]["template"])
            self.assertEqual(preview.rows, [])
            self.assertEqual(preview.validation["status"], "mapping_required")

    def test_inspect_workbook_preserves_normalized_export_validation_order(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "validation-order.xlsx"
            write_workbook(path, "early")

            preview = inspect_workbook(path)

            self.assertEqual(list(preview.validation), ["status", "counts", "errors"])

    def test_normalized_payload_preserves_legacy_replay_input_envelope(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "legacy-envelope.xlsx"
            write_workbook(path, "early")
            preview = inspect_workbook(path)

            payload = normalized_payload(preview)

            self.assertEqual(list(payload), [
                "schema_version", "workbook_filename", "source_sha256", "rows", "excluded_rows", "validation",
            ])
            self.assertNotIn("reporting_year", payload["rows"][0])
            self.assertEqual(payload["rows"][0]["field_provenance"], preview.rows[0]["field_provenance"])

    def test_inspect_workbook_preserves_original_whitespace_and_multiline_text(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "exact-text.xlsx"
            write_workbook(path, "early")
            workbook = load_workbook(path)
            workbook.active["C6"] = "  first line\n second line  "
            workbook.save(path)
            workbook.close()

            preview = inspect_workbook(path)

            self.assertEqual(preview.rows[0]["actual"], "  first line\n second line  ")
            self.assertEqual(preview.rows[0]["field_provenance"]["actual"]["original_value"], "  first line\n second line  ")
            self.assertEqual(list(Path(temp).iterdir()), [path])


if __name__ == "__main__":
    unittest.main()
