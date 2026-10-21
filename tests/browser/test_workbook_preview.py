import json
import re
import shutil
import sqlite3
from pathlib import Path

import httpx
from openpyxl import Workbook
from playwright.sync_api import Page, expect


def technical_workbook(path: Path, *, project: str = "SMART PARKING") -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BC Tuần 01.09 - 05.09"
    for address, value in {
        "A4": "STT", "B4": "Sản phẩm", "C4": "Kết quả tuần",
        "D4": "Kế hoạch tuần", "E4": "Vấn đề / rủi ro", "F4": "PIC",
        "A5": "1", "B5": project, "A6": "1.1", "B6": "PTSP",
        "C6": "Reported status from source", "D6": "Deliver next step", "F6": "Owner",
    }.items():
        sheet[address] = value
    workbook.save(path)
    workbook.close()
    return path


def mapping_required_workbook(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BC Tuần chưa xác định"
    sheet.append(["Mã nội bộ", "Nguồn dự án", "Nguồn vai trò", "Nguồn thực hiện", "Nguồn kế hoạch"])
    sheet.append(["1", "SMART PARKING", "PTSP", "Giá trị nguồn", "Bước tiếp theo"])
    workbook.save(path)
    workbook.close()
    return path


def multi_period_workbook(path: Path) -> Path:
    workbook = Workbook()
    sheets = [workbook.active, workbook.create_sheet("BC Tuần 31.08 - 05.09")]
    sheets[0].title = "BC Tuần 07.09 - 12.09"
    for sheet, actual in zip(sheets, ("  Bằng chứng kỳ mới\nDòng hai  ", "Bằng chứng kỳ cũ"), strict=True):
        for address, value in {
            "A4": "STT", "B4": "Sản phẩm", "C4": "Kết quả tuần",
            "D4": "Kế hoạch tuần", "E4": "Vấn đề / rủi ro", "F4": "PIC",
            "A5": "1", "B5": "SMART PARKING", "A6": "1.1", "B6": "PTSP",
            "C6": actual, "D6": "Bước tiếp theo", "F6": "Người phụ trách",
        }.items():
            sheet[address] = value
    workbook.save(path)
    workbook.close()
    return path


def source_status_workbook(path: Path) -> Path:
    workbook = Workbook()
    sheets = [
        workbook.active,
        workbook.create_sheet("Bố cục chưa nhận diện"),
        workbook.create_sheet("BC Tuần 31.08 - 05.09"),
        workbook.create_sheet("BC Tuần 15.08 - 20.08"),
    ]
    sheets[0].title = "BC Tuần 07.09 - 12.09"
    for address, value in {
        "A4": "STT", "B4": "Sản phẩm", "C4": "Kết quả tuần",
        "D4": "Kế hoạch tuần", "E4": "Vấn đề / rủi ro", "F4": "PIC",
        "A5": "1", "B5": "SMART PARKING", "A6": "1.1", "B6": "PTSP",
        "C6": "Exact source", "D6": "Next step", "F6": "Owner",
    }.items():
        sheets[0][address] = value
        sheets[2][address] = value
        sheets[3][address] = value
    for address, value in {"A1": "Mã", "B1": "Tùy chỉnh", "A2": "1", "B2": "Giá trị gốc"}.items():
        sheets[1][address] = value
    workbook.save(path)
    workbook.close()
    return path


def project_comparison_workbook(path: Path) -> Path:
    workbook = Workbook()
    current = workbook.active
    current.title = "BC Tuần 07.09 - 12.09"
    unmapped = workbook.create_sheet("Bố cục chưa nhận diện")
    previous = workbook.create_sheet("BC Tuần 31.08 - 05.09")
    earlier = workbook.create_sheet("BC Tuần 15.08 - 20.08")

    def populate(sheet, smart_actual: str | None, anat_actual: str, *, spaced_actual: str | None = None):
        for address, value in {
            "A4": "STT", "B4": "Sản phẩm", "C4": "Kết quả tuần",
            "D4": "Kế hoạch tuần", "E4": "Vấn đề / rủi ro", "F4": "PIC",
            "A5": "1", "B5": "SMART PARKING", "A6": "1.1", "B6": "PTSP",
            "A7": "2", "B7": "ANAT GSM", "A8": "2.1", "B8": "QC",
            "C8": anat_actual, "D8": "ANAT-PLAN-" + sheet.title, "F8": "ANAT-PIC",
        }.items():
            sheet[address] = value
        if smart_actual is not None:
            sheet["C6"] = smart_actual
            sheet["D6"] = "SMART-PLAN-" + sheet.title
            sheet["F6"] = "SMART-PIC"
        if spaced_actual is not None:
            sheet["A9"] = "3"
            sheet["B9"] = "SMART PARKING "
            sheet["A10"] = "3.1"
            sheet["B10"] = "PTSP"
            sheet["C10"] = spaced_actual
            sheet["D10"] = "SPACED-PLAN"
            sheet["F10"] = "SPACED-PIC"

    populate(current, "SMART-CURRENT-ONLY", "ANAT-CURRENT-ONLY")
    populate(previous, "SMART-PREVIOUS-ONLY", "ANAT-PREVIOUS-ONLY")
    populate(earlier, None, "ANAT-EARLIEST-ONLY", spaced_actual="SPACED-NAME-ONLY")
    unmapped["A1"] = "Custom project"
    unmapped["A2"] = "RAW-ONLY-PROJECT"
    workbook.save(path)
    workbook.close()
    return path


def test_catalog_identity(page: Page, app_server, app_url: str, external_attempts, tmp_path: Path):
    page.goto(f"{app_url}/app/")
    expect(page.get_by_role("heading", name="Chưa có workbook đã lưu")).to_be_visible()
    workbook = technical_workbook(tmp_path / "report.xlsx")
    with page.expect_response(lambda response: response.url.endswith("/api/workbooks") and response.request.method == "POST") as upload_info:
        page.get_by_label("Chọn tệp Excel").set_input_files(workbook)
        page.get_by_role("button", name="Tải và xem trước").click()
    upload = upload_info.value
    assert upload.status == 201
    workbook_id = upload.json()["workbook_id"]
    expect(page.get_by_role("heading", name="report.xlsx")).to_be_visible()

    changed = technical_workbook(tmp_path / "report.xlsx", project="ANAT GSM")
    page.goto(f"{app_url}/app/")
    with page.expect_response(lambda response: response.url.endswith("/api/workbooks") and response.request.method == "POST") as second_info:
        page.get_by_label("Chọn tệp Excel").set_input_files(changed)
        page.get_by_role("button", name="Tải và xem trước").click()
    second = second_info.value
    assert second.status == 201
    second_id = second.json()["workbook_id"]
    assert second_id != workbook_id

    page.goto(f"{app_url}/app/")
    row = page.get_by_role("row").filter(has_text="report.xlsx")
    expect(row).to_have_count(2)
    expect(page.get_by_role("heading", name="report.xlsx")).to_have_count(0)
    second_item = next(item for item in page.request.get(f"{app_url}/api/workbooks").json() if item["workbook_id"] == second_id)
    expect(row.first.locator("time")).to_have_text(re.compile(r"\d{2}/\d{2}/\d{4} \d{2}:\d{2} UTC"))
    page.get_by_role("link", name=f"Mở workbook report.xlsx, phiên bản {second_item['sha256'][:12]}").click()
    expect(page.get_by_role("heading", name="report.xlsx")).to_be_visible()
    assert page.url.endswith(f"/app/workbooks/{second_id}")


def test_dashboard_shell_keeps_unimplemented_steps_unavailable(page: Page, app_server, app_url: str, external_attempts):
    page.goto(f"{app_url}/app/")
    navigation = page.get_by_role("navigation", name="Tiến trình")
    expect(navigation).to_be_visible()
    expect(navigation.get_by_role("link", name="Tải báo cáo")).to_have_attribute("aria-current", "step")
    for step in ("Chọn hai kỳ", "Xem kết quả", "Review nhận xét", "Xuất kết quả"):
        expect(navigation.get_by_role("button", name=re.compile(step))).to_be_disabled()
    expect(navigation.get_by_text("Chưa khả dụng", exact=True)).to_have_count(4)
    expect(page.get_by_role("banner")).to_contain_text("Local · offline")
    expect(page.get_by_role("heading", name="Bắt đầu từ báo cáo của bạn.")).to_be_visible()
    expect(page.get_by_text("Bạn sẽ làm gì?", exact=True)).to_be_visible()
    expect(page.get_by_role("heading", name="Ít thao tác. Rõ việc cần làm.")).to_be_visible()
    expect(page.get_by_role("button", name="Dùng workbook minh họa")).to_have_count(0)
    page.set_viewport_size({"width": 320, "height": 900})
    assert page.evaluate("document.documentElement.scrollWidth <= document.documentElement.clientWidth")


def test_project_picker_filters_exact_project_across_multiple_sheets(page: Page, app_server, app_url: str, external_attempts, tmp_path: Path):
    page.goto(f"{app_url}/app/")
    source = project_comparison_workbook(tmp_path / "projects.xlsx")
    with page.expect_response(lambda response: response.url.endswith("/api/workbooks") and response.request.method == "POST"):
        page.get_by_label("Chọn tệp Excel").set_input_files(source)
        page.get_by_role("button", name="Tải và xem trước").click()
    expect(page.get_by_role("heading", name="projects.xlsx")).to_be_visible()

    project_picker = page.get_by_role("region", name="Chọn một dự án trong workbook")
    expect(project_picker).to_be_visible()
    expect(page.get_by_text("SMART-CURRENT-ONLY", exact=True)).to_have_count(0)
    expect(project_picker.get_by_role("button", name="Chọn dự án SMART PARKING", exact=True)).to_be_visible()
    expect(project_picker.get_by_role("button", name="Chọn dự án ANAT GSM", exact=True)).to_be_visible()
    expect(project_picker.get_by_role("button", name=re.compile("SMART PARKING.*khoảng trắng cuối"))).to_be_visible()
    expect(page.get_by_text("RAW-ONLY-PROJECT", exact=True)).to_have_count(0)
    expect(page.get_by_text("chưa thể xác định dự án nào có trong sheet", exact=False)).to_be_visible()
    search = project_picker.get_by_label("Tìm dự án trong workbook")
    search.fill("ANAT")
    expect(project_picker.get_by_role("button", name="Chọn dự án ANAT GSM", exact=True)).to_be_visible()
    expect(project_picker.get_by_role("button", name="Chọn dự án SMART PARKING", exact=True)).to_have_count(0)
    search.fill("")

    project_picker.get_by_role("button", name="Chọn dự án SMART PARKING", exact=True).click()
    expect(page.get_by_role("heading", name="SMART PARKING", exact=True)).to_be_visible()
    expect(page.get_by_text("SMART-CURRENT-ONLY", exact=True)).to_be_visible()
    expect(page.get_by_text("SMART-PREVIOUS-ONLY", exact=True)).to_be_visible()
    expect(page.get_by_text("ANAT-CURRENT-ONLY", exact=True)).to_have_count(0)
    expect(page.get_by_text("SPACED-NAME-ONLY", exact=True)).to_have_count(0)
    expect(page.get_by_text("RAW-ONLY-PROJECT", exact=True)).to_have_count(0)
    expect(page.get_by_role("button", name="Bỏ chọn kỳ BC Tuần 31.08 - 05.09")).to_have_attribute("aria-pressed", "true")

    page.get_by_role("button", name="Bỏ chọn kỳ BC Tuần 31.08 - 05.09").click()
    expect(page.get_by_text("SMART-PREVIOUS-ONLY", exact=True)).to_have_count(0)
    expect(page.get_by_text("SMART-CURRENT-ONLY", exact=True)).to_be_visible()
    page.get_by_role("button", name="Chọn kỳ BC Tuần 31.08 - 05.09").click()
    page.get_by_role("button", name="Đổi dự án").click()
    project_picker.get_by_role("button", name="Chọn dự án ANAT GSM", exact=True).click()
    expect(page.get_by_text("ANAT-CURRENT-ONLY", exact=True)).to_be_visible()
    expect(page.get_by_text("ANAT-PREVIOUS-ONLY", exact=True)).to_be_visible()
    expect(page.get_by_text("SMART-CURRENT-ONLY", exact=True)).to_have_count(0)

    page.get_by_role("button", name="Đổi dự án").click()
    spaced_option = project_picker.get_by_role("button", name=re.compile("SMART PARKING.*khoảng trắng cuối"))
    spaced_option.click()
    expect(page.get_by_text("SPACED-NAME-ONLY", exact=True)).to_be_visible()
    expect(page.get_by_text("SMART-CURRENT-ONLY", exact=True)).to_have_count(0)
    page.set_viewport_size({"width": 320, "height": 900})
    assert page.evaluate("document.documentElement.scrollWidth <= document.documentElement.clientWidth")
    period_comparison = page.get_by_role("region", name="Đối chiếu các kỳ đã chọn")
    assert period_comparison.evaluate("element => element.scrollWidth > element.clientWidth")
    page.get_by_role("button", name="Đổi dự án").click()
    project_cards = project_picker.locator(".project-list")
    assert project_cards.evaluate("element => element.scrollWidth > element.clientWidth")
    project_picker.get_by_role("button", name="Chọn dự án SMART PARKING", exact=True).click()
    expect(page.get_by_role("heading", name="SMART PARKING", exact=True)).to_be_visible()
    expect(page.get_by_text("ANAT-CURRENT-ONLY", exact=True)).to_have_count(0)
    page.get_by_role("button", name="Xem nguồn workbook").click()
    expect(page.get_by_role("heading", name="Các sheet trong nguồn")).to_be_visible()
    expect(page.get_by_role("list", name="Các sheet trong workbook")).to_be_visible()


def test_stale_selection(page: Page, app_server, app_url: str, external_attempts, tmp_path: Path):
    first_file = technical_workbook(tmp_path / "first.xlsx", project="SMART PARKING")
    second_file = technical_workbook(tmp_path / "second.xlsx", project="ANAT GSM")
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    first = page.request.post(f"{app_url}/api/workbooks", multipart={"workbook": {"name": "first.xlsx", "mimeType": mime, "buffer": first_file.read_bytes()}}).json()
    second = page.request.post(f"{app_url}/api/workbooks", multipart={"workbook": {"name": "second.xlsx", "mimeType": mime, "buffer": second_file.read_bytes()}}).json()
    page.add_init_script(f"""
      const originalFetch = window.fetch.bind(window);
      window.fetch = (input, init) => {{
        const url = typeof input === "string" ? input : input.url;
        if (url.endsWith("/api/workbooks/{first['workbook_id']}") && !window.__heldFirst) {{
          window.__heldFirst = true;
          return new Promise((resolve, reject) => {{
            window.__releaseFirst = () => originalFetch(input, init).then(resolve, reject);
          }});
        }}
        return originalFetch(input, init);
      }};
    """)
    page.goto(f"{app_url}/app/")
    page.get_by_role("link", name=f"Mở workbook first.xlsx, phiên bản {first['sha256'][:12]}").click()
    page.get_by_role("link", name=f"Mở workbook second.xlsx, phiên bản {second['sha256'][:12]}").click()
    expect(page.get_by_role("heading", name="second.xlsx")).to_be_visible()
    page.evaluate("window.__releaseFirst()")
    expect(page.get_by_role("heading", name="second.xlsx")).to_be_visible()
    assert page.url.endswith(f"/app/workbooks/{second['workbook_id']}")


def test_stale_upload_cannot_replace_newer_workbook_selection(page: Page, app_server, app_url: str, external_attempts, tmp_path: Path):
    existing = technical_workbook(tmp_path / "existing.xlsx", project="PROJECT-B")
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    existing_result = page.request.post(
        f"{app_url}/api/workbooks",
        multipart={"workbook": {"name": "existing.xlsx", "mimeType": mime, "buffer": existing.read_bytes()}},
    ).json()
    delayed_workbook_id = existing_result["workbook_id"] + 1
    page.add_init_script(f"""
      const originalFetch = window.fetch.bind(window);
      window.fetch = (input, init) => {{
        const url = typeof input === "string" ? input : input.url;
        if (url.endsWith("/api/workbooks/{delayed_workbook_id}") && !window.__heldUploadPreview) {{
          window.__heldUploadPreview = true;
          return new Promise((resolve, reject) => {{
            window.__releaseUploadPreview = () => originalFetch(input, init).then(resolve, reject);
          }});
        }}
        return originalFetch(input, init);
      }};
    """)
    page.goto(f"{app_url}/app/")
    new_workbook = technical_workbook(tmp_path / "new.xlsx", project="PROJECT-A")
    with page.expect_response(lambda response: response.url.endswith("/api/workbooks") and response.request.method == "POST") as upload_info:
        page.get_by_label("Chọn tệp Excel").set_input_files(new_workbook)
        page.get_by_role("button", name="Tải và xem trước").click()
    uploaded = upload_info.value
    assert uploaded.status == 201
    assert uploaded.json()["workbook_id"] == delayed_workbook_id

    page.get_by_role("link", name=f"Mở workbook existing.xlsx, phiên bản {existing_result['sha256'][:12]}").click()
    expect(page.get_by_role("heading", name="existing.xlsx")).to_be_visible()
    page.evaluate("window.__releaseUploadPreview()")
    expect(page.get_by_role("heading", name="existing.xlsx")).to_be_visible()
    assert page.url.endswith(f"/app/workbooks/{existing_result['workbook_id']}")


def test_upload_duplicate(page: Page, app_server, app_url: str, external_attempts, tmp_path: Path):
    page.goto(f"{app_url}/app/")
    first_file = technical_workbook(tmp_path / "first.xlsx")
    renamed_file = tmp_path / "renamed.xlsx"
    shutil.copyfile(first_file, renamed_file)
    posts = []
    page.on("request", lambda request: posts.append(request) if request.method == "POST" and request.url.endswith("/api/workbooks") else None)

    with page.expect_response(lambda response: response.url.endswith("/api/workbooks") and response.request.method == "POST") as first_info:
        page.get_by_label("Chọn tệp Excel").set_input_files(first_file)
        page.evaluate("document.querySelector('form').requestSubmit(); document.querySelector('form').requestSubmit();")
    first = first_info.value
    expect(page.get_by_role("heading", name="first.xlsx")).to_be_visible()
    first_data = first.json()
    assert first.status == 201
    assert len(posts) == 1
    catalog_item = page.request.get(f"{app_url}/api/workbooks").json()[0]
    imported_at = catalog_item["imported_at_utc"]

    page.goto(f"{app_url}/app/")
    with page.expect_response(lambda response: response.url.endswith("/api/workbooks") and response.request.method == "POST") as duplicate_info:
        page.get_by_label("Chọn tệp Excel").set_input_files(renamed_file)
        page.get_by_role("button", name="Tải và xem trước").click()
    duplicate = duplicate_info.value
    expect(page.locator("p[role='status']").first).to_contain_text("Tệp này đã có, đang mở bản đã lưu")
    assert duplicate.status == 200
    assert duplicate.json()["duplicate"] is True
    assert duplicate.json()["workbook_id"] == first_data["workbook_id"]
    after = page.request.get(f"{app_url}/api/workbooks").json()
    assert len(after) == 1
    assert after[0]["original_filename"] == "first.xlsx"
    assert after[0]["imported_at_utc"] == imported_at
    assert len(posts) == 2
    page.reload()
    assert len(posts) == 2

    with sqlite3.connect(app_server["database"]) as connection:
        assert connection.execute("SELECT COUNT(*) FROM workbook_versions").fetchone()[0] == 1
        assert connection.execute("SELECT COUNT(*) FROM analysis_jobs").fetchone()[0] == 0
        assert connection.execute("SELECT COUNT(*) FROM finding_reviews").fetchone()[0] == 0


def test_empty_source_states(page: Page, app_server, app_url: str, external_attempts, tmp_path: Path):
    with sqlite3.connect(app_server["database"]) as connection:
        connection.execute(
            "INSERT INTO workbook_versions(sha256, original_filename, storage_relpath, imported_at_utc) "
            "VALUES (?, ?, ?, ?)",
            ("f" * 64, "zero-sheets.xlsx", "ff/missing.xlsx", "2026-10-16T10:00:00+00:00"),
        )
    page.goto(f"{app_url}/app/")
    page.get_by_role("link", name="Mở workbook zero-sheets.xlsx").click()
    expect(page.get_by_text("Workbook chưa có sheet để xem", exact=False)).to_be_visible()

    unknown = mapping_required_workbook(tmp_path / "unknown.xlsx")
    with page.expect_response(lambda response: response.url.endswith("/api/workbooks") and response.request.method == "POST"):
        page.get_by_label("Chọn tệp Excel").set_input_files(unknown)
        page.get_by_role("button", name="Tải và xem trước").click()
    expect(page.get_by_role("heading", name="unknown.xlsx")).to_be_visible()
    expect(page.locator(".sheet-reason")).to_have_text("Cần thiết lập cột")
    page.get_by_role("button", name="Xem sheet BC Tuần chưa xác định").click()
    expect(page.get_by_text("số 0 không có nghĩa sheet Excel trống", exact=False)).to_be_visible()


def test_preview_overview(page: Page, app_server, app_url: str, external_attempts, tmp_path: Path):
    page.goto(f"{app_url}/app/")
    source = multi_period_workbook(tmp_path / "periods.xlsx")
    with page.expect_response(lambda response: response.url.endswith("/api/workbooks") and response.request.method == "POST"):
        page.get_by_label("Chọn tệp Excel").set_input_files(source)
        page.get_by_role("button", name="Tải và xem trước").click()
    expect(page.get_by_role("heading", name="periods.xlsx")).to_be_visible()
    attention = page.get_by_role("region", name="Mục cần chú ý từ nguồn")
    sheet_list = page.get_by_role("list", name="Các sheet trong workbook")
    expect(attention).to_be_visible()
    expect(attention).to_contain_text("Đây không phải kết quả phân tích chất lượng")
    expect(page.get_by_label("Tìm tên sheet")).to_be_visible()
    expect(page.get_by_role("group", name="Lọc sheet")).to_be_visible()
    items = sheet_list.get_by_role("listitem")
    expect(items).to_have_count(2)
    expect(items.nth(0)).to_contain_text("BC Tuần 07.09 - 12.09")
    expect(items.nth(1)).to_contain_text("BC Tuần 31.08 - 05.09")
    expect(page.get_by_role("heading", name="Chi tiết sheet")).to_have_count(0)
    page.get_by_role("button", name="Xem sheet BC Tuần 07.09 - 12.09").click()
    expect(page.get_by_role("heading", name="Chi tiết sheet: BC Tuần 07.09 - 12.09")).to_be_visible()
    expect(page.get_by_role("heading", name="Chi tiết sheet: BC Tuần 07.09 - 12.09")).to_be_focused()
    assert "?sheet=" in page.url
    expect(page.locator("[data-testid='selected-sheet-detail']")).to_have_count(1)
    page.set_viewport_size({"width": 320, "height": 900})
    overflow = page.evaluate("""() => ({
      viewport: window.innerWidth,
      document: document.documentElement.scrollWidth,
      body: document.body.scrollWidth,
          elements: [...document.body.querySelectorAll('*')].filter((element) => {
                const bounds = element.getBoundingClientRect();
                const parentBounds = element.parentElement?.getBoundingClientRect();
                return (bounds.right > window.innerWidth + 1 || bounds.left < -1) && !element.closest('.table-scroll')
                  && (!parentBounds || parentBounds.right <= window.innerWidth + 1);
              }).map((element) => ({tag: element.tagName, id: element.id, className: String(element.className), width: element.getBoundingClientRect().width, right: element.getBoundingClientRect().right, html: element.outerHTML.slice(0, 100)}))
    })""")
    assert overflow["document"] <= overflow["viewport"], overflow
    assert sheet_list.evaluate("element => element.scrollWidth <= element.clientWidth")


def test_sheet_search_filters_in_excel_order_and_attention_reveals_hidden_sheet(page: Page, app_server, app_url: str, external_attempts, tmp_path: Path):
    page.goto(f"{app_url}/app/")
    source = source_status_workbook(tmp_path / "search.xlsx")
    with page.expect_response(lambda response: response.url.endswith("/api/workbooks") and response.request.method == "POST") as upload_info:
        page.get_by_label("Chọn tệp Excel").set_input_files(source)
        page.get_by_role("button", name="Tải và xem trước").click()
    workbook_id = upload_info.value.json()["workbook_id"]
    expect(page.get_by_role("heading", name=source.name)).to_be_visible()
    expect(page).to_have_url(re.compile(rf".*/app/workbooks/{workbook_id}$"))
    data = page.request.get(f"{app_url}/api/workbooks/{workbook_id}").json()
    first, unmapped, conflict, year = data["sheets"]
    year_unknown = {"status": "identified", "label": "15.08-20.08", "year": None, "start_month_day": "08-15", "conflicts": []}
    conflict_period = {"status": "conflict", "label": "15.08-20.08", "year": None, "start_month_day": "08-15", "conflicts": ["Sheet/title period conflicts"]}
    with sqlite3.connect(app_server["database"]) as connection:
            connection.execute("UPDATE report_snapshots SET mapping_status='confirmed', sheet_state='draft_or_incomplete', reporting_period_json=? WHERE snapshot_id=?", (json.dumps(year_unknown), first["snapshot_id"]))
            connection.execute("UPDATE report_snapshots SET mapping_status='recognized', sheet_state='not_assessed', reporting_period_json=? WHERE snapshot_id=?", (json.dumps(conflict_period), conflict["snapshot_id"]))
            connection.execute("UPDATE report_snapshots SET mapping_status='recognized', sheet_state='not_assessed', reporting_period_json=? WHERE snapshot_id=?", (json.dumps(year_unknown), year["snapshot_id"]))
    page.reload()

    attention = page.get_by_role("region", name="Mục cần chú ý từ nguồn")
    expect(attention.get_by_role("listitem")).to_have_count(3)
    expect(attention).to_contain_text("Thông tin kỳ mâu thuẫn")
    expect(attention).to_contain_text("Nháp hoặc chưa đầy đủ")
    year_context = page.get_by_role("region", name="Thông tin năm còn thiếu")
    expect(year_context).to_contain_text("1 sheet chưa rõ năm")
    year_context.get_by_text("1 sheet chưa rõ năm", exact=True).click()
    expect(year_context.get_by_role("listitem")).to_have_count(1)
    expect(year_context).to_contain_text("Chưa rõ năm")
    expect(page.get_by_text("Vì sao?", exact=True)).to_be_visible()
    expect(page.get_by_role("button", name=re.compile("Đối chiếu nguồn Bố cục chưa nhận diện"))).to_be_visible()

    sheet_list = page.get_by_role("list", name="Các sheet trong workbook")
    items = sheet_list.get_by_role("listitem")
    expect(items).to_have_count(4)
    expect(items.nth(1)).to_contain_text("Bố cục chưa nhận diện")
    page.get_by_role("button", name="Cần chú ý · 3").click()
    expect(items).to_have_count(3)
    expect(items.nth(0)).to_contain_text("BC Tuần 07.09 - 12.09")
    page.get_by_role("button", name="Tất cả · 4").click()
    expect(items).to_have_count(4)
    search = page.get_by_label("Tìm tên sheet")
    search.fill("không tồn tại")
    expect(page.get_by_role("status", name="Số sheet đang hiển thị")).to_contain_text("0 / 4")
    expect(page.get_by_text("Không tìm thấy sheet. Thử tên khác", exact=False)).to_be_visible()
    search.fill("BC Tuần")
    expect(items).to_have_count(3)
    expect(items.nth(0)).to_contain_text("BC Tuần 07.09 - 12.09")
    page.get_by_role("button", name=re.compile("Chưa rõ năm · 1")).click()
    expect(page.get_by_role("button", name=re.compile("Chưa rõ năm · 1"))).to_have_attribute("aria-pressed", "true")
    search.fill("không tồn tại")
    page.get_by_role("button", name="Xóa tìm kiếm tên sheet").click()
    expect(search).to_have_value("")
    expect(page.get_by_role("button", name=re.compile("Chưa rõ năm · 1"))).to_have_attribute("aria-pressed", "true")
    expect(items).to_have_count(1)
    search.fill("không tồn tại")
    page.get_by_role("button", name=re.compile("Đối chiếu nguồn Bố cục chưa nhận diện")).click()
    detail_heading = page.get_by_role("heading", name="Chi tiết sheet: Bố cục chưa nhận diện")
    expect(detail_heading).to_be_focused()
    page.get_by_role("button", name=re.compile("Đối chiếu nguồn Bố cục chưa nhận diện")).click()
    expect(detail_heading).to_be_focused()
    expect(page.get_by_role("button", name=re.compile("Tất cả · 4"))).to_have_attribute("aria-pressed", "true")
    assert page.url.endswith(f"?sheet={unmapped['snapshot_id']}")
    page.set_viewport_size({"width": 320, "height": 900})
    assert page.evaluate("document.documentElement.scrollWidth <= document.documentElement.clientWidth")
    for selector in ("#sheet-search-1", ".sheet-filters button", ".sheet-row button", ".sheet-metadata summary"):
        control = page.locator(selector).first
        assert control.evaluate("element => element.getBoundingClientRect().height >= 44")
    page.locator("#sheet-search-1").focus()
    page.keyboard.press("Tab")
    assert page.evaluate("document.activeElement.matches(':focus-visible') && getComputedStyle(document.activeElement).outlineStyle !== 'none'")
    page.locator(".sheet-filters button").nth(1).focus()
    page.keyboard.press("Enter")
    expect(page.locator(".sheet-filters button").nth(1)).to_have_attribute("aria-pressed", "true")
    page.locator(".sheet-row button").first.focus()
    page.keyboard.press("Enter")
    page.locator(".sheet-row").first.locator(".sheet-metadata summary").focus()
    page.keyboard.press("Space")


def test_attention_counts_are_dto_driven_and_generic_warnings_stay_secondary(page: Page, app_server, app_url: str, external_attempts, tmp_path: Path):
    page.goto(f"{app_url}/app/")
    source = source_status_workbook(tmp_path / "dynamic-attention.xlsx")
    with page.expect_response(lambda response: response.url.endswith("/api/workbooks") and response.request.method == "POST") as upload_info:
        page.get_by_label("Chọn tệp Excel").set_input_files(source)
        page.get_by_role("button", name="Tải và xem trước").click()
    workbook_id = upload_info.value.json()["workbook_id"]
    expect(page.get_by_role("heading", name=source.name)).to_be_visible()
    expect(page).to_have_url(re.compile(rf".*/app/workbooks/{workbook_id}$"))
    data = page.request.get(f"{app_url}/api/workbooks/{workbook_id}").json()
    sheets = data["sheets"]
    normal_period = {"status": "identified", "label": "31.08-05.09/2026", "year": 2026, "start_month_day": "08-31", "conflicts": []}
    unknown_period = {"status": "unknown", "label": None, "year": None, "start_month_day": None, "conflicts": []}
    year_period = {"status": "identified", "label": "15.08-20.08", "year": None, "start_month_day": "08-15", "conflicts": []}
    generic_warning = ["<b>Unclassified source warning</b>"]
    with sqlite3.connect(app_server["database"]) as connection:
        connection.execute("UPDATE report_snapshots SET mapping_status='recognized', sheet_state='not_assessed', reporting_period_json=?, warnings_json=?", (json.dumps(normal_period), json.dumps(generic_warning)))
    page.reload()
    attention = page.get_by_role("region", name="Mục cần chú ý từ nguồn")
    expect(attention).to_contain_text("Chưa ghi nhận mục cần chú ý")
    expect(attention.get_by_role("listitem")).to_have_count(0)
    expect(page.get_by_role("button", name="Tất cả · 4")).to_be_visible()
    expect(page.get_by_role("button", name="Cần chú ý · 0")).to_be_visible()
    expect(page.get_by_role("button", name="Chưa rõ năm · 0")).to_be_visible()
    first_row = page.locator(".sheet-list > li").first
    first_row.get_by_text("Thông tin sheet", exact=True).click()
    expect(first_row).to_contain_text("<b>Unclassified source warning</b>")
    assert page.locator("b").count() == 0

    with sqlite3.connect(app_server["database"]) as connection:
        connection.execute("UPDATE report_snapshots SET sheet_state='draft_or_incomplete' WHERE snapshot_id=?", (sheets[0]["snapshot_id"],))
        connection.execute("UPDATE report_snapshots SET reporting_period_json=? WHERE snapshot_id=?", (json.dumps(unknown_period), sheets[1]["snapshot_id"]))
        connection.execute("UPDATE report_snapshots SET reporting_period_json=? WHERE snapshot_id IN (?, ?)", (json.dumps(year_period), sheets[2]["snapshot_id"], sheets[3]["snapshot_id"]))
    page.reload()
    attention = page.get_by_role("region", name="Mục cần chú ý từ nguồn")
    expect(attention.get_by_role("listitem")).to_have_count(2)
    expect(attention).to_contain_text("Chưa xác định kỳ")
    expect(attention).to_contain_text("Nháp hoặc chưa đầy đủ")
    expect(attention).not_to_contain_text("Unclassified source warning")
    expect(page.get_by_role("button", name="Cần chú ý · 2")).to_be_visible()
    year_region = page.get_by_role("region", name="Thông tin năm còn thiếu")
    expect(year_region).to_contain_text("2 sheet chưa rõ năm")


def test_source_disclosure(page: Page, app_server, app_url: str, external_attempts, tmp_path: Path):
    page.goto(f"{app_url}/app/")
    source = multi_period_workbook(tmp_path / "periods.xlsx")
    with page.expect_response(lambda response: response.url.endswith("/api/workbooks") and response.request.method == "POST"):
        page.get_by_label("Chọn tệp Excel").set_input_files(source)
        page.get_by_role("button", name="Tải và xem trước").click()
    page.get_by_role("button", name="Xem sheet BC Tuần 07.09 - 12.09").click()
    disclosure = page.get_by_text("Dòng nguồn và vị trí ô", exact=True)
    details = disclosure.locator("xpath=..")
    assert details.evaluate("element => element.open") is False
    disclosure.click()
    source_table = page.get_by_role("table", name="Dữ liệu nguồn gốc")
    expect(source_table).to_contain_text("Bằng chứng kỳ mới")
    expect(source_table).to_contain_text("C6")
    expect(source_table).to_contain_text("Kết quả tuần")
    assert source_table.locator("td[data-field='actual'] .source-value").text_content() == "  Bằng chứng kỳ mới\nDòng hai  "


def test_preview_status_warnings(page: Page, app_server, app_url: str, external_attempts, tmp_path: Path):
    page.goto(f"{app_url}/app/")
    source = source_status_workbook(tmp_path / "statuses.xlsx")
    with page.expect_response(lambda response: response.url.endswith("/api/workbooks") and response.request.method == "POST") as upload_info:
        page.get_by_label("Chọn tệp Excel").set_input_files(source)
        page.get_by_role("button", name="Tải và xem trước").click()
    workbook_id = upload_info.value.json()["workbook_id"]
    expect(page.get_by_role("heading", name="statuses.xlsx")).to_be_visible()
    assert page.url.endswith(f"/app/workbooks/{workbook_id}")
    details = page.request.get(f"{app_url}/api/workbooks/{workbook_id}").json()
    sheets = details["sheets"]
    first_period = {"status": "identified", "label": "07.09-12.09", "year": None, "start_month_day": "09-07", "conflicts": []}
    certain_period = {"status": "identified", "label": "31.08-05.09/2026", "year": 2026, "start_month_day": "08-31", "conflicts": []}
    conflict_period = {"status": "conflict", "label": "15.08-20.08", "year": None, "start_month_day": "08-15", "conflicts": ["Sheet/title period conflicts"]}
    with sqlite3.connect(app_server["database"]) as connection:
        connection.execute(
                "UPDATE report_snapshots SET mapping_status=?, sheet_state=?, reporting_period_json=?, warnings_json=? WHERE snapshot_id=?",
                ("confirmed", "draft_or_incomplete", json.dumps(first_period), json.dumps([
                "Project and role share a physical header; use the recorded hierarchy",
                "Column D can contain a more-specific role without an independent header",
                "Explicit flat mapping; blank project/role labels are not inferred",
                "<img src=x onerror=alert(1)>",
            ]), sheets[0]["snapshot_id"]),
        )
        connection.execute(
                "UPDATE report_snapshots SET mapping_status=?, sheet_state=?, reporting_period_json=?, warnings_json=? WHERE snapshot_id=?",
            ("mapping_required", "mapping_required", "null", json.dumps(["Unknown worksheet layout; mapping is required"]), sheets[1]["snapshot_id"]),
        )
        connection.execute(
            "UPDATE report_snapshots SET mapping_status=?, sheet_state=?, reporting_period_json=?, warnings_json=? WHERE snapshot_id=?",
            ("recognized", "not_assessed", json.dumps(certain_period), "[]", sheets[2]["snapshot_id"]),
        )
        connection.execute(
            "UPDATE report_snapshots SET mapping_status=?, sheet_state=?, reporting_period_json=?, warnings_json=? WHERE snapshot_id=?",
            ("recognized", "not_assessed", json.dumps(conflict_period), "[]", sheets[3]["snapshot_id"]),
        )

    page.reload()
    attention = page.get_by_role("region", name="Mục cần chú ý từ nguồn")
    expect(attention.get_by_role("listitem")).to_have_count(3)
    expect(attention).to_contain_text("Cần thiết lập cột")
    expect(attention).to_contain_text("Thông tin kỳ mâu thuẫn")
    expect(attention).to_contain_text("Nháp hoặc chưa đầy đủ")
    year_context = page.get_by_role("region", name="Thông tin năm còn thiếu")
    expect(year_context.get_by_role("listitem")).to_have_count(0)
    sheet_list = page.get_by_role("list", name="Các sheet trong workbook")
    rows = page.locator(".sheet-list > li")
    expect(rows.nth(0)).to_contain_text("4 cảnh báo nguồn")
    expect(rows.nth(1)).to_contain_text("Cần thiết lập cột")
    expect(rows.nth(2)).to_contain_text("Đã xác định kỳ và năm")
    expect(rows.nth(3)).to_contain_text("Thông tin kỳ mâu thuẫn")
    rows.nth(0).get_by_text("Thông tin sheet", exact=True).click()
    expect(rows.nth(0)).to_contain_text("confirmed")
    expect(rows.nth(0)).to_contain_text("draft_or_incomplete")
    expect(rows.nth(0)).to_contain_text("Chưa rõ năm")
    expect(rows.nth(0)).to_contain_text("<img src=x onerror=alert(1)>")
    expect(rows.nth(1)).to_contain_text("Unknown worksheet layout; mapping is required")
    rows.nth(0).get_by_text("Thông tin sheet", exact=True).click()
    rows.nth(1).get_by_text("Thông tin sheet", exact=True).click()

    page.get_by_role("button", name="Xem sheet BC Tuần 07.09 - 12.09").click()
    details_disclosure = page.get_by_test_id("selected-sheet-detail").locator("details.technical-details > summary")
    details_disclosure.click()
    expect(page.get_by_text("confirmed", exact=True)).to_be_visible()
    expect(page.get_by_text("draft_or_incomplete", exact=True)).to_be_visible()
    expect(page.get_by_role("heading", name="Cột dự án và vai trò dùng chung cấu trúc")).to_be_visible()
    expect(page.get_by_role("heading", name="Cột D có thể chứa vai trò chi tiết hơn")).to_be_visible()
    expect(page.get_by_role("heading", name="Nhãn trống không được tự điền")).to_be_visible()
    raw_warning = "<img src=x onerror=alert(1)>"
    expect(page.get_by_test_id("selected-sheet-detail").locator(".warning-list pre").filter(has_text=raw_warning)).not_to_be_visible()
    page.get_by_text("Cảnh báo gốc", exact=True).last.click()
    expect(page.get_by_test_id("selected-sheet-detail").locator(".warning-list pre").filter(has_text=raw_warning)).to_be_visible()
    assert page.locator("img[src='x']").count() == 0


def test_upload_tracer(page: Page, app_server, app_url: str, external_attempts, tmp_path: Path):
    page.goto(f"{app_url}/app/")
    file_input = page.get_by_label("Chọn tệp Excel")
    file_input.set_input_files(technical_workbook(tmp_path / "technical.xlsx"))
    with page.expect_response(
        lambda response: "/api/workbooks/" in response.url and response.request.method == "GET"
    ) as response_info:
        page.get_by_role("button", name="Tải và xem trước").click()
    preview_response = response_info.value
    assert preview_response.status == 200
    preview = preview_response.json()
    assert preview["sheets"][0]["sheet_name"] == "BC Tuần 01.09 - 05.09"
    expect(page.get_by_text("BC Tuần 01.09 - 05.09", exact=True)).to_be_visible()
    expect(page.get_by_role("heading", name="technical.xlsx")).to_be_focused()
    page.get_by_role("button", name="Xem sheet BC Tuần 01.09 - 05.09").click()
    page.get_by_text("Dòng nguồn và vị trí ô", exact=True).click()
    expect(page.get_by_text("Reported status from source", exact=True)).to_be_visible()
    expect(page.locator("p.step-note")).to_contain_text("Đã lưu workbook, chưa chạy phân tích.")
    await_empty = page.locator("#workbook-file").input_value()
    assert await_empty == ""
    assert page.url.endswith(f"/app/workbooks/{preview['workbook_id']}?sheet={preview['sheets'][0]['snapshot_id']}")

    with sqlite3.connect(app_server["database"]) as connection:
        assert connection.execute("SELECT COUNT(*) FROM workbook_versions").fetchone()[0] == 1
        assert connection.execute("SELECT COUNT(*) FROM analysis_jobs").fetchone()[0] == 0


def test_upload_rejection_api(app_server, tmp_path: Path):
    base_url = app_server["base_url"]
    with httpx.Client(base_url=base_url) as client:
        missing = client.post("/api/workbooks")
        assert missing.status_code == 422

        rejected = [
            ("empty.xlsx", b""),
            ("corrupt.xlsx", b"not an xlsx archive"),
            ("report.csv", technical_workbook(tmp_path / "wrong-extension.xlsx").read_bytes()),
            ("large.xlsx", b"x" * (32 * 1024 * 1024 + 1)),
        ]
        for filename, content in rejected:
            response = client.post(
                "/api/workbooks",
                files={"workbook": (filename, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
            assert response.status_code == 400, filename
            assert str(tmp_path) not in response.text

        foreign_origin = client.post(
            "/api/workbooks",
            files={"workbook": ("foreign.xlsx", technical_workbook(tmp_path / "foreign.xlsx").read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers={"Origin": "https://evil.example"},
        )
        assert foreign_origin.status_code == 403

    with sqlite3.connect(app_server["database"]) as connection:
        assert connection.execute("SELECT COUNT(*) FROM workbook_versions").fetchone()[0] == 0
        assert connection.execute("SELECT COUNT(*) FROM analysis_jobs").fetchone()[0] == 0


def test_upload_rejection_preserves_file_and_does_not_retry(page: Page, app_server, app_url: str, external_attempts, tmp_path: Path):
    page.goto(f"{app_url}/app/")
    selected_file = technical_workbook(tmp_path / "bad.xlsx")
    selected_file.write_bytes(b"invalid workbook bytes")
    input_control = page.get_by_label("Chọn tệp Excel")
    input_control.set_input_files(selected_file)
    post_attempts = []
    page.on("request", lambda request: post_attempts.append(request) if request.method == "POST" and request.url.endswith("/api/workbooks") else None)

    def non_json_response(route):
        route.fulfill(status=200, content_type="text/html", body="<html>not json</html>")

    page.route("**/api/workbooks", non_json_response)
    page.get_by_role("button", name="Tải và xem trước").click()
    expect(page.get_by_role("status")).to_contain_text("phản hồi không phải JSON")
    assert page.url.endswith("/app/")
    assert page.locator("#workbook-file").evaluate("element => element.files[0].name") == "bad.xlsx"
    assert page.locator(".workbook-preview").count() == 0

    page.unroute("**/api/workbooks", non_json_response)
    page.get_by_role("button", name="Tải và xem trước").click()
    expect(page.get_by_role("status")).to_contain_text("Không thể đọc workbook")
    assert page.url.endswith("/app/")
    assert page.locator("#workbook-file").evaluate("element => element.files[0].name") == "bad.xlsx"
    assert page.locator(".workbook-preview").count() == 0

    assert len(post_attempts) == 2
    page.reload()
    assert len(post_attempts) == 2
    with sqlite3.connect(app_server["database"]) as connection:
        assert connection.execute("SELECT COUNT(*) FROM workbook_versions").fetchone()[0] == 0
        assert connection.execute("SELECT COUNT(*) FROM analysis_jobs").fetchone()[0] == 0

    input_control.set_input_files(technical_workbook(tmp_path / "corrected.xlsx"))
    page.get_by_role("button", name="Tải và xem trước").click()
    expect(page.get_by_text("BC Tuần 01.09 - 05.09", exact=True)).to_be_visible()
    assert page.url.endswith("/app/workbooks/1")
