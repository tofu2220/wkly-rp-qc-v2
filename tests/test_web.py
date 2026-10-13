import asyncio
import sqlite3
import tempfile
import unittest
from pathlib import Path

import httpx
from openpyxl import Workbook

from wkly_rp_qc_v2.store import (
    create_job,
    get_snapshot,
    get_template_mapping,
    resolve_project_alias,
)
from wkly_rp_qc_v2.web import WebSettings, create_app
from wkly_rp_qc_v2.web_service import confirm_mapping


def workbook_bytes(*, unknown: bool = False, project: str = "SMART PARKING") -> bytes:
    with tempfile.TemporaryDirectory() as temp:
        path = Path(temp) / "upload.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Unknown Layout" if unknown else "BC Tuần 01.09 - 05.09"
        if unknown:
            rows = [
                ["Project", "Role", "PIC", "Actual", "Next Plan", "Issue/Risk"],
                [project, "PTSP", "Owner", "Reported status", "Deliver next step", None],
                [None, None, None, "Unlabeled continuation", None, None],
            ]
            for row_index, row in enumerate(rows, start=1):
                for column, value in enumerate(row, start=1):
                    if value is not None:
                        sheet.cell(row_index, column, value)
        else:
            for address, value in {
                "A4": "STT", "B4": "Sản phẩm", "C4": "Kết quả tuần",
                "D4": "Kế hoạch tuần", "E4": "Vấn đề / rủi ro", "F4": "PIC",
                "A5": "1", "B5": project, "A6": "1.1", "B6": "PTSP",
                "C6": "Reported status", "D6": "Deliver next step", "F6": "Owner",
            }.items():
                sheet[address] = value
        workbook.save(path)
        workbook.close()
        return path.read_bytes()


def two_period_workbook_bytes(project="ANAT GSM") -> bytes:
    with tempfile.TemporaryDirectory() as temp:
        path = Path(temp) / "two-period.xlsx"
        workbook = Workbook()
        for index, (title, actual, next_plan) in enumerate((
            ("BC Tuần 31.08 - 05.09", "Previous status", "Deliver API"),
            ("BC Tuần 07.09 - 12.09", "Current update", "Next task"),
        )):
            sheet = workbook.active if index == 0 else workbook.create_sheet()
            sheet.title = title
            for address, value in {
                "A4": "STT", "B4": "Sản phẩm", "C4": "Kết quả tuần",
                "D4": "Kế hoạch tuần", "E4": "Vấn đề / rủi ro", "F4": "PIC",
                "A5": "1", "B5": project, "A6": "1.1", "B6": "PTSP",
                "C6": actual, "D6": next_plan, "F6": "Owner",
            }.items():
                sheet[address] = value
        workbook.save(path)
        workbook.close()
        return path.read_bytes()


class WebAppTests(unittest.TestCase):
    def make_client(self, root: Path):
        settings = WebSettings(
            database_path=root / "app.sqlite3",
            storage_dir=root / "storage",
            replay_roots=(),
        )
        return ASGITestClient(create_app(settings))

    def upload(self, client, content: bytes, filename="report.xlsx"):
        return client.post(
            "/workbooks",
            files={"workbook": (filename, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            follow_redirects=False,
        )

    def test_home_page_renders_local_workflow(self):
        with tempfile.TemporaryDirectory() as temp, self.make_client(Path(temp)) as client:
            response = client.get("/")

            self.assertEqual(response.status_code, 200)
            self.assertIn("Workbooks", response.text)
            self.assertIn("127.0.0.1", response.text)
            self.assertIn("Replay offline", response.text)

    def test_upload_previews_known_layout_without_scheduling_a_job(self):
        with tempfile.TemporaryDirectory() as temp, self.make_client(Path(temp)) as client:
            response = self.upload(client, workbook_bytes())
            self.assertEqual(response.status_code, 303)
            workbook_id = int(response.headers["location"].rsplit("/", 1)[1])
            preview = client.get(response.headers["location"])

            self.assertEqual(preview.status_code, 200)
            self.assertIn("BC Tuần 01.09 - 05.09", preview.text)
            self.assertIn("recognized", preview.text)
            self.assertIn("Reported status", preview.text)
            self.assertIn("C6", preview.text)
            with sqlite3.connect(Path(temp) / "app.sqlite3") as connection:
                self.assertEqual(connection.execute("SELECT COUNT(*) FROM analysis_jobs").fetchone()[0], 0)
                self.assertEqual(connection.execute("SELECT COUNT(*) FROM workbook_versions").fetchone()[0], 1)
            snapshot = get_snapshot(Path(temp) / "app.sqlite3", 1)
            self.assertEqual(snapshot["workbook_id"], workbook_id)

    def test_duplicate_upload_reuses_existing_workbook_version(self):
        with tempfile.TemporaryDirectory() as temp, self.make_client(Path(temp)) as client:
            content = workbook_bytes()
            first = self.upload(client, content)
            second = self.upload(client, content)

            self.assertEqual(first.status_code, 303)
            self.assertEqual(first.headers["location"], second.headers["location"])
            with sqlite3.connect(Path(temp) / "app.sqlite3") as connection:
                self.assertEqual(connection.execute("SELECT COUNT(*) FROM workbook_versions").fetchone()[0], 1)

    def test_unknown_layout_requires_explicit_flat_mapping_and_does_not_fill_labels(self):
        with tempfile.TemporaryDirectory() as temp, self.make_client(Path(temp)) as client:
            response = self.upload(client, workbook_bytes(unknown=True), "unknown.xlsx")
            workbook_id = int(response.headers["location"].rsplit("/", 1)[1])
            preview = client.get(response.headers["location"])
            self.assertIn("mapping_required", preview.text)
            self.assertIn("Save mapping", preview.text)
            self.assertIn("Reported status", preview.text)
            self.assertIn("Unlabeled continuation", preview.text)

            mapped = client.post(
                f"/workbooks/{workbook_id}/sheets/Unknown%20Layout/mapping",
                data={
                    "template_id": "manual-flat", "version": "1", "source_reference": "PM column mapping",
                    "header_row": "1", "first_data_row": "2", "project_col": "1", "role_col": "2",
                    "pic_col": "3", "actual_col": "4", "next_plan_col": "5", "issue_risk_col": "6",
                },
                follow_redirects=False,
            )

            self.assertEqual(mapped.status_code, 303)
            preview = client.get(mapped.headers["location"])
            self.assertIn("confirmed", preview.text)
            self.assertIn("manual-flat@1", preview.text)
            self.assertIn("PM column mapping", preview.text)
            snapshot = get_snapshot(Path(temp) / "app.sqlite3", 1)
            self.assertEqual(snapshot["mapping_status"], "confirmed")
            self.assertEqual(snapshot["rows"][0]["project_label"], "SMART PARKING")
            self.assertIsNone(snapshot["rows"][1]["project_label"])
            self.assertIsNone(snapshot["rows"][1]["role_label"])
            self.assertIn("manual_mapping_missing_project_label", snapshot["rows"][1]["review_flags"])
            self.assertFalse(resolve_project_alias(Path(temp) / "app.sqlite3", "SMART PARKING")["explicit"])

    def test_alias_is_only_used_after_explicit_confirmation(self):
        with tempfile.TemporaryDirectory() as temp, self.make_client(Path(temp)) as client:
            response = self.upload(client, workbook_bytes(project="ANAT GSM"))
            workbook_id = int(response.headers["location"].rsplit("/", 1)[1])
            before = resolve_project_alias(Path(temp) / "app.sqlite3", "ANAT GSM")
            mapped = client.post(
                f"/workbooks/{workbook_id}/aliases",
                data={"raw_label": "ANAT GSM", "canonical_label": "DEPOT GSM", "source_reference": "explicit user decision"},
                follow_redirects=False,
            )
            after = resolve_project_alias(Path(temp) / "app.sqlite3", "ANAT GSM")
            snapshot = get_snapshot(Path(temp) / "app.sqlite3", 1)

            self.assertIsNone(before["project_id"])
            self.assertEqual(mapped.status_code, 303)
            self.assertTrue(after["explicit"])
            self.assertEqual(after["canonical_label"], "DEPOT GSM")
            self.assertEqual(snapshot["rows"][0]["project_label"], "ANAT GSM")

    def test_scope_uses_only_registered_project_aliases(self):
        with tempfile.TemporaryDirectory() as temp, self.make_client(Path(temp)) as client:
            response = self.upload(client, two_period_workbook_bytes())
            workbook_id = int(response.headers["location"].rsplit("/", 1)[1])
            client.post(
                f"/workbooks/{workbook_id}/aliases",
                data={"raw_label": "ANAT GSM", "canonical_label": "DEPOT GSM", "source_reference": "explicit project mapping"},
            )

            unaliased = client.get(
                "/scope", params={
                    "previous_snapshot_id": 1, "current_snapshot_id": 2,
                    "project": "DEPOT GSM", "role": "PTSP",
                },
            )
            aliased = client.get(
                "/scope", params={
                    "previous_snapshot_id": 1, "current_snapshot_id": 2,
                    "project": "DEPOT GSM", "role": "PTSP", "aliases": ["ANAT GSM"],
                },
            )

            self.assertIn("No previous-week plans", unaliased.text)
            self.assertIn("Previous rows: 1", aliased.text)
            self.assertIn("ANAT GSM", aliased.text)

    def test_mapping_decision_is_not_saved_if_analysis_already_uses_snapshot(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            settings = WebSettings(root / "store.sqlite3", root / "storage", ())
            app = create_app(settings)
            with ASGITestClient(app) as client:
                response = self.upload(client, workbook_bytes(unknown=True), "unknown.xlsx")
            workbook_id = int(response.headers["location"].rsplit("/", 1)[1])
            database = settings.database_path
            with sqlite3.connect(database) as connection:
                current_id = connection.execute(
                    "INSERT INTO report_snapshots(workbook_id, sheet_name, sheet_index, template_id, "
                    "mapping_status, sheet_state, reporting_period_json, mapping_json, warnings_json, row_count) "
                    "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    (workbook_id, "current", 1, "early", "recognized", "not_assessed", "{}", "{}", "[]", 0),
                ).lastrowid
            create_job(
                database, workbook_id=workbook_id, previous_snapshot_id=1, current_snapshot_id=current_id,
                project_label="SMART PARKING", role="PTSP", mode="replay", replay_bundle_id="bundle-1",
            )
            mapping = {
                "mode": "flat", "template_id": "manual-flat", "version": "2",
                "header_row": 1, "first_data_row": 2,
                "columns": {"project": 1, "role": 2, "pic": 3, "actual": 4, "next_plan": 5, "issue_risk": 6},
            }

            with self.assertRaisesRegex(ValueError, "after an analysis job"):
                confirm_mapping(settings, workbook_id, "Unknown Layout", mapping, "explicit remapping test")

            self.assertIsNone(get_template_mapping(database, "manual-flat", "2"))


if __name__ == "__main__":
    unittest.main()


class ASGITestClient:
    def __init__(self, app):
        self.app = app

    def __enter__(self):
        return self

    def __exit__(self, *_):
        return None

    async def _request(self, method, path, **kwargs):
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=self.app), base_url="http://testserver"
        ) as client:
            return await client.request(method, path, **kwargs)

    def get(self, path, **kwargs):
        return asyncio.run(self._request("GET", path, **kwargs))

    def post(self, path, **kwargs):
        return asyncio.run(self._request("POST", path, **kwargs))
