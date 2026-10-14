import csv
import io
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from .report import _cell, _display

JsonDict = dict[str, Any]
EXPORT_FIELDS = (
    "project", "role", "reporting_period", "source_sha256", "plan_id", "planned_commitment",
    "previous_period_reference", "current_reported_text", "source_reported_status", "warning_type",
    "model_severity", "model_severity_status", "effective_severity", "effective_severity_status",
    "approval_state", "policy_decision_id", "source_reference", "finding", "missing_information",
    "question_for_PM", "review_state", "review_note", "source_locations", "source_quotes",
    "closure_status", "milestone_change",
)


def _locations(evidence: list[JsonDict], sources: JsonDict) -> list[str]:
    locations = []
    for item in evidence:
        source = sources.get(item["source_id"])
        if source is not None:
            locations.append(f"{source['sheet']}!{source['cell']} (row {source['source_row']}, {source['field']})")
    return list(dict.fromkeys(locations))


def _export_rows(comparison: JsonDict, reviews: list[JsonDict]) -> list[JsonDict]:
    scope = comparison["scope"]
    metadata = comparison.get("metadata", {})
    sources = comparison["sources"]
    commitments = {item["plan_id"]: item for item in comparison["commitments"]}
    policy = {
        (record["plan_id"], record["finding_index"]): record
        for record in metadata.get("policy_resolution", [])
    }
    review_by_finding = {
        (record["plan_id"], record["finding_index"]): record
        for record in reviews
    }
    rows = []
    for assessment in comparison["assessments"]:
        plan_id = assessment["plan_id"]
        commitment = commitments[plan_id]
        plan_evidence = commitment["plan_evidence"]
        update_evidence = assessment["update_evidence"]
        for finding_index, finding in enumerate(assessment["findings"]):
            resolution = policy.get((plan_id, finding_index))
            review = review_by_finding.get((plan_id, finding_index), {})
            all_evidence = plan_evidence + update_evidence + finding["evidence"]
            model_severity = resolution["model_severity"] if resolution else finding["severity"]
            model_status = resolution["model_severity_status"] if resolution else finding["severity_status"]
            rows.append({
                "project": scope["project"],
                "role": metadata.get("role", ""),
                "reporting_period": f"{scope['previous_sheet']} → {scope['current_sheet']}",
                "source_sha256": metadata.get("input_sha256", ""),
                "plan_id": plan_id,
                "planned_commitment": commitment["planned_commitment"],
                "previous_period_reference": json.dumps(plan_evidence, ensure_ascii=False),
                "current_reported_text": json.dumps(update_evidence, ensure_ascii=False),
                "source_reported_status": assessment["source_reported_status"] or "",
                "warning_type": finding["rule"],
                "model_severity": model_severity or "",
                "model_severity_status": model_status,
                "effective_severity": finding["severity"] or "",
                "effective_severity_status": finding["severity_status"],
                "approval_state": resolution["effective_severity_status"] if resolution else "unavailable",
                "policy_decision_id": (
                    resolution["decision_id"] if resolution and resolution["decision_id"] else
                    ("no approved decision" if resolution else "not recorded")
                ),
                "source_reference": resolution["source_reference"] if resolution and resolution["source_reference"] else "",
                "finding": finding["explanation"],
                # The strict assessment schema has no separate missing_information field.
                "missing_information": "",
                "question_for_PM": finding["question_for_PM"],
                "review_state": review.get("state", "not_reviewed"),
                "review_note": review.get("note") or "",
                "source_locations": json.dumps(_locations(all_evidence, sources), ensure_ascii=False),
                "source_quotes": json.dumps(all_evidence, ensure_ascii=False),
                "closure_status": None,
                "milestone_change": None,
            })
    return rows


def _markdown(comparison: JsonDict, rows: list[JsonDict]) -> str:
    metadata = comparison.get("metadata", {})
    lines = [
        "# Rà soát chất lượng báo cáo tuần", "",
        f"- Dự án: {_cell(comparison['scope']['project'])}",
        f"- Kỳ: {_cell(rows[0]['reporting_period']) if rows else '—'}",
        f"- Cam kết: {comparison['counts']['commitments']}; có phát hiện: {comparison['counts']['affected_commitments']}; findings: {comparison['counts']['findings']}.",
        "- Trạng thái PM được lưu riêng; không thay đổi finding, evidence hoặc policy approval.",
        "", "## Chính sách và nguồn quyết định",
    ]
    if "policy_resolution" not in metadata:
        lines.append("Policy provenance unavailable; no approval is inferred for this legacy result.")
    else:
        lines.extend([
            f"- Policy: {_cell(metadata.get('policy_version', 'unavailable'))}",
            f"- SHA-256: {_cell(metadata.get('policy_sha256', 'unavailable'))}",
        ])
    for row in rows:
        lines.extend([
            "", f"## { _cell(row['plan_id']) } · {_cell(row['warning_type'])}",
            f"- Cam kết: {_cell(row['planned_commitment'])}",
            f"- Previous-period reference: {_cell(row['previous_period_reference']) or '—'}",
            f"- Current reported status: {_cell(row['source_reported_status']) or 'Không được báo cáo'}",
            f"- Mức hiệu lực: {_cell(row['effective_severity'] or 'Chưa xác định')} ({_cell(row['effective_severity_status'])})",
            f"- Model: {_cell(row['model_severity'] or 'Chưa xác định')} ({_cell(row['model_severity_status'])})",
            f"- Quyết định: {_cell(row['policy_decision_id']) or 'no approved decision'} · nguồn: {_cell(row['source_reference']) or '—'}",
            f"- Cập nhật nguồn: {_cell(row['current_reported_text']) or '—'}",
            f"- Finding: {_cell(row['finding'])}",
            f"- Thông tin còn thiếu: {_cell(row['missing_information']) or 'Chưa được cấu trúc riêng trong assessment'}",
            f"- Câu hỏi PM: {_cell(row['question_for_PM'])}",
            f"- Review PM: {_cell(row['review_state'])} · {_cell(row['review_note']) or '—'}",
            f"- Nguồn: {_cell(row['source_locations']) or '—'}",
            f"- Closure status: {_cell(row['closure_status']) or 'unknown'} · milestone change: {_cell(row['milestone_change']) or 'unknown'}",
        ])
    return "\n".join(lines) + "\n"


def write_web_exports(run_dir: Path, comparison: JsonDict, reviews: list[JsonDict]) -> None:
    rows = _export_rows(comparison, reviews)
    csv_stream = io.StringIO(newline="")
    writer = csv.DictWriter(csv_stream, fieldnames=EXPORT_FIELDS, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow({key: _display(value) for key, value in row.items()})
    (run_dir / "pm_review.csv").write_text("\ufeff" + csv_stream.getvalue(), encoding="utf-8", newline="")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Findings"
    sheet.append(EXPORT_FIELDS)
    for row in rows:
        cells = []
        for field in EXPORT_FIELDS:
            value = row[field]
            cells.append("" if value is None else value if isinstance(value, (int, float)) else str(value))
        sheet.append(cells)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    coverage = workbook.create_sheet("Coverage")
    coverage.append(("Metric", "Value"))
    for field, value in comparison["counts"].items():
        coverage.append((field, value))
    workbook.save(run_dir / "pm_review.xlsx")
    workbook.close()
    (run_dir / "pm_review.md").write_text(_markdown(comparison, rows), encoding="utf-8")
