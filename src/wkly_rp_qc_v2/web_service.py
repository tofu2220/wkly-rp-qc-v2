import tempfile
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .core import JsonDict, PreparedScope, prepare_scope
from .importer import inspect_workbook
from .store import (
    ImportResult,
    apply_template_mapping,
    confirm_project_alias,
    get_original_workbook_path,
    get_snapshot,
    import_workbook,
    list_project_aliases,
    list_snapshots,
)

MAX_UPLOAD_BYTES = 32 * 1024 * 1024


@dataclass(frozen=True)
class WebSettings:
    database_path: Path
    storage_dir: Path
    replay_roots: tuple[Path, ...] = ()


def import_upload(settings: WebSettings, filename: str | None, content: bytes) -> ImportResult:
    safe_name = Path((filename or "upload.xlsx").replace("\\", "/")).name or "upload.xlsx"
    if Path(safe_name).suffix.casefold() != ".xlsx":
        raise ValueError("Only .xlsx workbooks are supported")
    if len(content) > MAX_UPLOAD_BYTES:
        raise ValueError("Workbook exceeds the 32 MiB local upload limit")
    settings.storage_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="upload-", dir=settings.storage_dir) as staging:
        source = Path(staging) / safe_name
        source.write_bytes(content)
        return import_workbook(source, settings.database_path, settings.storage_dir)


def workbook_details(settings: WebSettings, workbook_id: int) -> dict:
    snapshots = [item for item in list_snapshots(settings.database_path) if item["workbook_id"] == workbook_id]
    if not snapshots:
        return {}
    sheets = [get_snapshot(settings.database_path, item["snapshot_id"]) for item in snapshots]
    if any(sheet["mapping_status"] == "mapping_required" for sheet in sheets):
        original = get_original_workbook_path(settings.database_path, settings.storage_dir, workbook_id)
        workbook = load_workbook(original, data_only=False, read_only=True)
        try:
            for sheet in sheets:
                if sheet["mapping_status"] != "mapping_required":
                    continue
                source_sheet = workbook[sheet["sheet_name"]]
                preview_rows = []
                for row_number in range(1, min(source_sheet.max_row, 20) + 1):
                    cells = []
                    for column in range(1, min(source_sheet.max_column, 12) + 1):
                        cell = source_sheet.cell(row_number, column)
                        value = str(cell.value) if cell.value is not None else None
                        truncated = bool(value and len(value) > 240)
                        cells.append({
                            "address": f"{get_column_letter(column)}{row_number}",
                            "column": get_column_letter(column),
                            "value": value[:240] + ("…" if truncated else "") if value is not None else None,
                            "truncated": truncated,
                        })
                    preview_rows.append({"row": row_number, "cells": cells})
                sheet["cell_preview"] = preview_rows
        finally:
            workbook.close()
    return {
        "workbook_id": workbook_id,
        "sha256": snapshots[0]["workbook_sha256"],
        "original_filename": snapshots[0]["original_filename"],
        "sheets": sheets,
    }


def confirm_mapping(
    settings: WebSettings,
    workbook_id: int,
    sheet_name: str,
    mapping: JsonDict,
    source_reference: str,
) -> int:
    original = get_original_workbook_path(settings.database_path, settings.storage_dir, workbook_id)
    preview = inspect_workbook(original, mapping_overrides={sheet_name: mapping})
    sheet = next((item for item in preview.sheets if item["sheet"] == sheet_name), None)
    if sheet is None or sheet["mapping_status"] != "confirmed" or preview.validation["errors"]:
        raise ValueError("The explicit sheet mapping did not validate")
    return apply_template_mapping(
        settings.database_path, settings.storage_dir, workbook_id, sheet_name,
        mapping["template_id"], mapping["version"], mapping,
        source_reference=source_reference,
    )


def confirm_alias(settings: WebSettings, raw_label: str, canonical_label: str, source_reference: str) -> int:
    return confirm_project_alias(
        settings.database_path, raw_label, canonical_label, source_reference=source_reference
    )


def scope_preview(
    settings: WebSettings,
    *,
    previous_snapshot_id: int,
    current_snapshot_id: int,
    project: str,
    role: str,
    aliases: list[str] | None = None,
) -> PreparedScope:
    registered_aliases = set(list_project_aliases(settings.database_path, project))
    if any(alias not in registered_aliases for alias in aliases or []):
        raise ValueError("Project alias must be explicitly confirmed for the selected canonical label")
    snapshots = (
        get_snapshot(settings.database_path, previous_snapshot_id),
        get_snapshot(settings.database_path, current_snapshot_id),
    )
    if any(snapshot is None for snapshot in snapshots):
        raise ValueError("Selected report snapshot is unavailable")
    rows = []
    for snapshot in snapshots:
        for row in snapshot["rows"]:
            rows.append({
                "sheet": snapshot["sheet_name"],
                "source_row": row["source_row"],
                "project": row["project_label"],
                "role": row["role_label"],
                "pic": row["pic"],
                "actual": row["actual"],
                "next_plan": row["next_plan"],
                "issue_risk": row["issue_risk"],
                "source_refs": row["source_refs"],
                "review_flags": row["review_flags"],
                "reporting_period": row["reporting_period"],
                "reporting_year": row["reporting_year"],
            })
    return prepare_scope(
        rows,
        previous=snapshots[0]["sheet_name"],
        current=snapshots[1]["sheet_name"],
        project=project,
        role=role,
        aliases=aliases or [],
    )
