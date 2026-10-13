import hashlib
import re
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter

from .core import JsonDict

FIELDS = ("project", "role", "pic", "actual", "next_plan", "issue_risk")
SHEET_PERIOD_RE = re.compile(r"(\d{1,2})\.(\d{1,2})\s*-\s*(\d{1,2})\.(\d{1,2})")
TITLE_PERIOD_RE = re.compile(r"(\d{1,2})/(\d{1,2})\s*[–-]\s*(\d{1,2})/(\d{1,2})(?:/(\d{4}))?")
NUMBER_ID_RE = re.compile(r"^\d+(?:\.\d+)*\.?$")
ROMAN_ID_RE = re.compile(r"^[IVXLCDM]+\.?$", re.IGNORECASE)
TOTAL_RE = re.compile(r"^\s*(?:tổng cộng|tổng số|grand total|total)\s*$", re.IGNORECASE)


@dataclass(frozen=True)
class WorkbookImportPreview:
    workbook_sha256: str
    workbook_filename: str
    sheets: list[JsonDict]
    rows: list[JsonDict]
    excluded_rows: JsonDict
    validation: JsonDict


def _serializable(value: Any) -> tuple[Any, str]:
    if value is None:
        return None, "null"
    if isinstance(value, datetime):
        return value.isoformat(timespec="microseconds"), "datetime_iso8601"
    if isinstance(value, date):
        return value.isoformat(), "date_iso8601"
    if isinstance(value, time):
        return value.isoformat(timespec="microseconds"), "time_iso8601"
    if isinstance(value, timedelta):
        return value.total_seconds(), "duration_seconds"
    if isinstance(value, (str, int, float, bool)):
        return value, type(value).__name__
    return str(value), "rich_text_or_other_as_displayed_string"


def _value_at(sheet, row: int, column: int) -> Any:
    cell = sheet.cell(row, column)
    return None if isinstance(cell, MergedCell) else cell.value


def _text_at(sheet, row: int, column: int) -> str | None:
    value = _value_at(sheet, row, column)
    return str(value) if value is not None else None


def _meaningful(value: Any) -> bool:
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def _merge_index(sheet) -> dict[tuple[int, int], tuple[str, str, int, int]]:
    by_row_col = {}
    for span in sheet.merged_cells.ranges:
        name = str(span)
        anchor = sheet.cell(span.min_row, span.min_col).coordinate
        for row in range(span.min_row, span.max_row + 1):
            for column in range(span.min_col, span.max_col + 1):
                by_row_col[(row, column)] = (anchor, name, span.min_row, span.min_col)
    return by_row_col


def _field_from_cell(sheet, row: int, column: int, merges, *, resolve_merged_label=False) -> JsonDict:
    coordinate = f"{get_column_letter(column)}{row}"
    original, kind = _serializable(_value_at(sheet, row, column))
    info = merges.get((row, column))
    anchor, merged_range, anchor_row, anchor_column = info if info else (coordinate, None, row, column)
    value = original
    resolution = "direct" if original is not None else "blank"
    if info and (anchor_row != row or anchor_column != column) and resolve_merged_label:
        value, _ = _serializable(_value_at(sheet, anchor_row, anchor_column))
        resolution = "merged_anchor" if value is not None else "blank_merged_anchor"
    return {
        "source_cell": coordinate,
        "original_value": original,
        "resolved_value": value,
        "original_value_kind": kind,
        "anchor_cell": anchor,
        "anchor_original_value": _serializable(_value_at(sheet, anchor_row, anchor_column))[0],
        "merged_range": merged_range,
        "resolution": resolution,
    }


def _contextual_field(sheet, row: int, column: int, label_row: int, label_column: int, merges) -> JsonDict:
    result = _field_from_cell(sheet, row, column, merges)
    label = _field_from_cell(sheet, label_row, label_column, merges, resolve_merged_label=True)
    result.update(
        resolved_value=label["resolved_value"],
        anchor_cell=label["anchor_cell"],
        merged_range=label["merged_range"],
        resolution="hierarchical_project_group",
        context_cell=label["source_cell"],
        context_original_value=label["original_value"],
        anchor_original_value=label["anchor_original_value"],
    )
    result["original_value"] = None
    result["original_value_kind"] = "null"
    return result


def _header_has(sheet, row: int, column: int, fragment: str) -> bool:
    value = _text_at(sheet, row, column)
    return bool(value and fragment.casefold() in value.casefold())


def _get_template(sheet) -> JsonDict | None:
    if _text_at(sheet, 4, 1) and _text_at(sheet, 4, 1).strip().casefold() == "stt":
        if all((
            _header_has(sheet, 4, 2, "sản phẩm"),
            _header_has(sheet, 4, 3, "kết quả"),
            _header_has(sheet, 4, 4, "kế hoạch"),
            _header_has(sheet, 4, 5, "vấn đề"),
            _header_has(sheet, 4, 6, "pic"),
        )):
            return {
                "name": "early", "header_row": 4, "id_col": 1, "label_col": 2,
                "subrole_col": None, "project_col": 2, "role_col": 2,
                "actual_col": 3, "next_plan_col": 4, "issue_risk_col": 5, "pic_col": 6,
            }
        return None

    if _header_has(sheet, 5, 3, "vai trò") and all((
        _header_has(sheet, 5, 5, "kết quả"),
        _header_has(sheet, 5, 6, "kế hoạch"),
        _header_has(sheet, 5, 7, "vấn đề"),
        _header_has(sheet, 5, 8, "pic"),
    )):
        return {
            "name": "revised", "header_row": 5, "id_col": 2, "label_col": 3,
            "subrole_col": 4, "project_col": 3, "role_col": 3,
            "actual_col": 5, "next_plan_col": 6, "issue_risk_col": 7, "pic_col": 8,
        }
    if _header_has(sheet, 5, 4, "vai trò") and all((
        _header_has(sheet, 5, 6, "kết quả"),
        _header_has(sheet, 5, 7, "kế hoạch"),
        _header_has(sheet, 5, 8, "vấn đề"),
        _header_has(sheet, 5, 9, "pic"),
    )):
        return {
            "name": "latest", "header_row": 5, "id_col": 3, "project_id_col": 2,
            "label_col": 4, "subrole_col": None, "project_col": 4, "role_col": 4,
            "actual_col": 6, "next_plan_col": 7, "issue_risk_col": 8, "pic_col": 9,
        }
    return None


def _period_of(sheet, config) -> JsonDict:
    name_match = SHEET_PERIOD_RE.search(sheet.title)
    if not name_match:
        return {"status": "unknown", "label": None, "year": None, "source": None,
                "conflicts": ["No date range found in sheet name"]}
    start_day, start_month, end_day, end_month = map(int, name_match.groups())
    label = f"{start_day:02}.{start_month:02}-{end_day:02}.{end_month:02}"
    title_address = "B3" if config["header_row"] == 5 else "A2"
    title_text = _text_at(sheet, 3, 2) if config["header_row"] == 5 else _text_at(sheet, 2, 1)
    title_match = TITLE_PERIOD_RE.search(title_text or "")
    flags = []
    title_period = None
    year = None
    if title_match:
        day1, month1, day2, month2 = map(int, title_match.groups()[:4])
        title_period = f"{day1:02}.{month1:02}-{day2:02}.{month2:02}"
        if title_period != label:
            flags.append(f"Title {title_address} period {title_period} conflicts with sheet-name period {label}")
        elif title_match.group(5):
            year = int(title_match.group(5))
    else:
        flags.append(f"No unambiguous reporting period parsed from title {title_address}")
    actual_header = _text_at(sheet, config["header_row"], config["actual_col"])
    actual_match = SHEET_PERIOD_RE.search(actual_header or "")
    actual_period = None
    if actual_match:
        day1, month1, day2, month2 = map(int, actual_match.groups())
        actual_period = f"{day1:02}.{month1:02}-{day2:02}.{month2:02}"
        if actual_period != label:
            flags.append(f"Actual header period {actual_period} conflicts with sheet-name period {label}")
    return {
        "status": "conflict" if any("conflicts" in flag for flag in flags) else "identified",
        "label": label if year is None else f"{label}/{year}",
        "start_month_day": f"{start_month:02}-{start_day:02}",
        "end_month_day": f"{end_month:02}-{end_day:02}",
        "year": year,
        "source": "sheet_name; year only if matching explicit title",
        "title_cell": title_address,
        "title_text": title_text,
        "title_period": title_period,
        "actual_header_period": actual_period,
        "conflicts": flags,
    }


def _header_mapping(sheet, config, merges) -> JsonDict:
    row = config["header_row"]
    mapping = {}
    for field in FIELDS:
        column = config.get(f"{field}_col", config["label_col"])
        cell = sheet.cell(row, column)
        info = merges.get((row, column))
        if info and isinstance(cell, MergedCell):
            cell = sheet[info[0]]
        extra = ["D"] if field == "role" and config["name"] == "revised" else []
        mapping[field] = {
            "status": "contextual_hierarchy_no_explicit_project_header" if field == "project" else
                      ("verified_with_subrole_column" if extra else "verified"),
            "column_letter": get_column_letter(column),
            "additional_columns": extra,
            "original_header_text": cell.value,
            "header_cell_address": cell.coordinate,
            "header_merged_range": merges.get((row, column), (None, None))[1],
            "first_data_row": row + 1,
        }
    return mapping


def _manual_mapping(sheet, mapping: JsonDict) -> tuple[JsonDict, JsonDict]:
    if set(mapping) != {"mode", "template_id", "version", "header_row", "first_data_row", "columns"} or mapping["mode"] != "flat":
        raise ValueError("Invalid manual flat mapping fields")
    if not isinstance(mapping["template_id"], str) or not mapping["template_id"].strip():
        raise ValueError("Manual mapping template_id must be nonempty")
    if not isinstance(mapping["version"], str) or not mapping["version"].strip():
        raise ValueError("Manual mapping version must be nonempty")
    header_row, first_data_row = mapping["header_row"], mapping["first_data_row"]
    if (type(header_row) is not int or type(first_data_row) is not int
            or header_row < 1 or first_data_row <= header_row or first_data_row > sheet.max_row + 1):
        raise ValueError("Manual mapping requires valid header_row and first_data_row")
    columns = mapping["columns"]
    if not isinstance(columns, dict) or set(columns) != set(FIELDS):
        raise ValueError("Manual mapping columns must define all six normalized fields")
    for field, column in columns.items():
        if field in ("project", "role") and column is None:
            raise ValueError(f"Manual mapping requires a {field} column")
        if column is not None and (type(column) is not int or not 1 <= column <= sheet.max_column):
            raise ValueError(f"Invalid manual mapping column for {field}")
    if not any(columns[field] is not None for field in ("actual", "next_plan", "issue_risk")):
        raise ValueError("Manual mapping requires at least one narrative column")

    config = {
        "name": "manual_flat", "template_id": mapping["template_id"], "version": mapping["version"],
        "header_row": header_row, "first_data_row": first_data_row, "columns": columns,
        "actual_col": columns.get("actual") or columns.get("next_plan") or columns.get("issue_risk") or columns["project"],
    }
    fields = {}
    for field, column in columns.items():
        if column is None:
            fields[field] = {
                "status": "unmapped", "column_letter": None, "additional_columns": [],
                "original_header_text": None, "header_cell_address": None,
                "header_merged_range": None, "first_data_row": first_data_row,
            }
        else:
            address = f"{get_column_letter(column)}{header_row}"
            fields[field] = {
                "status": "user_mapped", "column_letter": get_column_letter(column), "additional_columns": [],
                "original_header_text": _value_at(sheet, header_row, column),
                "header_cell_address": address, "header_merged_range": None,
                "first_data_row": first_data_row,
            }
    return config, fields


def _manual_flat_rows(sheet, config) -> tuple[list[JsonDict], list[JsonDict]]:
    rows, excluded = [], [
        {"row": row, "reason": "header_or_title"} for row in range(1, config["first_data_row"])
    ]
    for row in range(config["first_data_row"], sheet.max_row + 1):
        fields, values, source_refs = {}, {}, {}
        for field, column in config["columns"].items():
            if column is None:
                fields[field] = {
                    "source_cell": None, "original_value": None, "resolved_value": None,
                    "original_value_kind": "null", "anchor_cell": None, "anchor_original_value": None,
                    "merged_range": None, "resolution": "unmapped",
                }
                values[field] = None
                source_refs[field] = None
            else:
                item = _field_from_cell(sheet, row, column, {})
                fields[field] = item
                values[field] = item["resolved_value"]
                source_refs[field] = item["source_cell"]
        if not any(_meaningful(value) for value in values.values()):
            excluded.append({"row": row, "reason": "unlabelled_empty_data_candidate"})
            continue
        review = []
        if not _meaningful(values["project"]):
            review.append("manual_mapping_missing_project_label")
        if not _meaningful(values["role"]):
            review.append("manual_mapping_missing_role_label")
        has_narrative = any(_meaningful(values[field]) for field in ("actual", "next_plan", "issue_risk"))
        if not has_narrative and not _meaningful(values["pic"]):
            review.append("draft_role_row_no_report_values")
        rows.append({
            "sheet": sheet.title, "source_row": row, **values, "source_refs": source_refs,
            "field_provenance": fields, "subgroup": [], "review_flags": review,
        })
    return rows, excluded


def _number_code(raw: Any) -> str | None:
    if raw is None:
        return None
    text = str(raw).strip().rstrip(".")
    return text if NUMBER_ID_RE.fullmatch(text) else None


def _code_parent(code: str) -> str | None:
    return code.rsplit(".", 1)[0] if "." in code else None


def _roman_number(value: str | None) -> int | None:
    if not value or not ROMAN_ID_RE.fullmatch(value.strip()):
        return None
    digits = {"I": 1, "V": 5, "X": 10, "L": 50, "C": 100, "D": 500, "M": 1000}
    chars = value.strip().upper().rstrip(".")
    return sum(-digits[c] if i + 1 < len(chars) and digits[c] < digits[chars[i + 1]] else digits[c]
               for i, c in enumerate(chars))


def _sheet_rows(sheet, config, merges) -> tuple[list[JsonDict], list[JsonDict]]:
    head = config["header_row"]
    label_column = config["label_col"]
    id_column = config["id_col"]
    columns = {field: config[f"{field}_col"] for field in FIELDS if field not in ("project", "role")}
    indices, row_codes, descendants = {}, {}, set()
    for row in range(head + 1, sheet.max_row + 1):
        code = _number_code(_text_at(sheet, row, id_column))
        if code:
            row_codes[row] = code
            indices.setdefault(code, row)
    for code in indices:
        parent = _code_parent(code)
        while parent:
            descendants.add(parent)
            parent = _code_parent(parent)

    groups, row_kind = {}, {}
    for row in range(head + 1, sheet.max_row + 1):
        primary = _text_at(sheet, row, label_column)
        supplement = _text_at(sheet, row, config["subrole_col"]) if config["subrole_col"] else None
        code = row_codes.get(row)
        has_report = any(_meaningful(_value_at(sheet, row, columns[field]))
                         for field in ("actual", "next_plan", "issue_risk"))
        if primary and TOTAL_RE.fullmatch(primary):
            row_kind[row] = "total"
        elif config["name"] == "latest" and _text_at(sheet, row, config["project_id_col"]) and ROMAN_ID_RE.fullmatch(_text_at(sheet, row, config["project_id_col"]).strip()):
            row_kind[row] = "project_group"
            groups[row] = {"code": None, "label_col": label_column, "kind": "project"}
        elif config["name"] == "revised" and _roman_number(_text_at(sheet, row, id_column)) is not None:
            is_category = str(primary or "").casefold().strip() == "an ninh an toàn"
            row_kind[row] = "section" if is_category else "roman_project_group"
            if not is_category:
                groups[row] = {"code": None, "label_col": label_column, "kind": "roman_project"}
        elif code and (code in descendants or (config["name"] == "early" and "." not in code)) and not has_report:
            row_kind[row] = "hierarchy_group"
            groups[row] = {"code": code, "label_col": label_column, "kind": "group"}
        elif code or has_report or _meaningful(_value_at(sheet, row, columns["pic"])) or _meaningful(supplement):
            row_kind[row] = "data_candidate"
        elif _meaningful(primary):
            row_kind[row] = "section"
        else:
            row_kind[row] = "blank"

    normalized, excluded = [], []
    last_project_row = last_roman_project_row = last_named_legacy_group_row = None
    for row in range(1, sheet.max_row + 1):
        if row <= head:
            excluded.append({"row": row, "reason": "header_or_title" if row != 3 else "title_or_spacer"})
            continue
        kind = row_kind.get(row, "blank")
        if kind == "project_group" and config["name"] == "latest":
            last_project_row = row
        if kind == "roman_project_group" and config["name"] == "revised":
            last_roman_project_row = row
        if kind == "hierarchy_group" and config["name"] == "early":
            group_code = row_codes[row]
            root_label_row = row
            ancestor = _code_parent(group_code)
            while ancestor:
                if indices.get(ancestor) in groups:
                    root_label_row = indices[ancestor]
                ancestor = _code_parent(ancestor)
            last_named_legacy_group_row = root_label_row
        if kind != "data_candidate":
            excluded.append({"row": row, "reason": kind})
            continue
        if config["name"] == "early" and row > head + 2:
            values = [_value_at(sheet, row, column) for column in (
                id_column, label_column, columns["actual"], columns["next_plan"], columns["issue_risk"], columns["pic"]
            )]
            prior = [any(_meaningful(_value_at(sheet, previous, column)) for column in (
                id_column, label_column, columns["actual"], columns["next_plan"], columns["issue_risk"], columns["pic"]
            )) for previous in range(max(head + 1, row - 3), row)]
            if (not any(_meaningful(value) for value in (values[0], values[1], values[3], values[4], values[5]))
                    and values[2] in ("x", "X") and not any(prior)):
                excluded.append({"row": row, "reason": "isolated_tail_marker_not_report_data"})
                continue

        has_content = any(_meaningful(_value_at(sheet, row, columns[field]))
                          for field in ("actual", "next_plan", "issue_risk"))
        code = row_codes.get(row)
        code_entry = merges.get((row, id_column))
        if not code and code_entry:
            code = _number_code(_text_at(sheet, code_entry[2], id_column))
        if not code and config["subrole_col"]:
            label_merge = merges.get((row, label_column))
            if label_merge:
                code = row_codes.get(label_merge[2])

        role = _field_from_cell(sheet, row, label_column, merges, resolve_merged_label=True)
        if config["subrole_col"]:
            additional = _field_from_cell(sheet, row, config["subrole_col"], merges, resolve_merged_label=True)
            if _meaningful(additional["resolved_value"]):
                additional["parent_role"] = role["resolved_value"]
                additional["parent_role_anchor_cell"] = role["anchor_cell"]
                role = additional

        review, project_row, subgroups = [], None, []
        if config["name"] == "latest":
            project_row = last_project_row
            if code:
                for prefix in [".".join(code.split(".")[:i]) for i in range(1, len(code.split(".")))]:
                    subgroup_row = indices.get(prefix)
                    if subgroup_row in groups and subgroup_row != project_row:
                        subgroups.append({"label": _text_at(sheet, subgroup_row, label_column),
                                          "source_cell": f"{get_column_letter(label_column)}{subgroup_row}"})
        else:
            if code:
                parent = _code_parent(code)
                ancestors = []
                while parent:
                    group_row = indices.get(parent)
                    if group_row is not None and group_row in groups:
                        ancestors.append(group_row)
                    parent = _code_parent(parent)
                if ancestors:
                    project_row = ancestors[-1]
                    for subgroup_row in reversed(ancestors[:-1]):
                        subgroups.append({"label": _text_at(sheet, subgroup_row, label_column),
                                          "source_cell": f"{get_column_letter(label_column)}{subgroup_row}"})
            if (config["name"] == "early" and project_row is not None and code
                    and last_named_legacy_group_row is not None and last_named_legacy_group_row > project_row):
                latest_code = row_codes.get(last_named_legacy_group_row)
                if latest_code and code != latest_code and not code.startswith(latest_code + "."):
                    review.append("stt_conflicts_with_nearest_project_heading")
                    project_row, subgroups = None, []
            if project_row is None and config["name"] == "revised" and last_roman_project_row:
                section_number = _roman_number(_text_at(sheet, last_roman_project_row, id_column))
                code_root = int(code.split(".")[0]) if code else None
                if section_number == code_root:
                    project_row = last_roman_project_row
                elif code is None:
                    project_row = last_roman_project_row
                    review.append("roman_group_used_for_unnumbered_continuation")
            if project_row is None and config["name"] == "early" and not code and last_named_legacy_group_row:
                project_row = last_named_legacy_group_row
                review.append("nonmerged_project_context_from_nearest_named_section_review")
            if project_row is None:
                review.append("project_context_uncertain_no_numbered_parent")

        if project_row is not None:
            project = _contextual_field(sheet, row, config["project_col"], project_row, label_column, merges)
        else:
            project = _field_from_cell(sheet, row, config["project_col"], merges)
            project.update(original_value=None, original_value_kind="null", resolved_value=None,
                           resolution="unresolved_group")
            if config["name"] == "early" and last_named_legacy_group_row:
                project["candidate_group_cell"] = f"{get_column_letter(label_column)}{last_named_legacy_group_row}"
                project["candidate_group_value"] = _text_at(sheet, last_named_legacy_group_row, label_column)
        if not _meaningful(role["resolved_value"]):
            review.append("role_missing_or_nonmerged_continuation")
        fields = {"project": project, "role": role, "pic": _field_from_cell(sheet, row, columns["pic"], merges)}
        for field in ("actual", "next_plan", "issue_risk"):
            fields[field] = _field_from_cell(sheet, row, columns[field], merges)
        source_refs = {
            field: (obj.get("context_cell") or obj["anchor_cell"] if field == "project" else obj["source_cell"])
            for field, obj in fields.items()
        }
        if not any((has_content, _meaningful(fields["role"]["resolved_value"]),
                    _meaningful(fields["pic"]["original_value"]))):
            excluded.append({"row": row, "reason": "unlabelled_empty_data_candidate"})
            continue
        if project_row is None:
            review.append("project_needs_human_review")
        if code is None and not fields["role"]["resolved_value"]:
            review.append("unlabelled_row_with_report_content_no_safe_role_inference")
        if not has_content and not _meaningful(fields["pic"]["original_value"]):
            review.append("draft_role_row_no_report_values")
        normalized.append({
            "sheet": sheet.title, "source_row": row,
            **{field: fields[field]["resolved_value"] for field in FIELDS},
            "source_refs": source_refs, "field_provenance": fields,
            "subgroup": subgroups, "review_flags": list(dict.fromkeys(review)),
        })
    return normalized, excluded


def _validate_extraction(workbook, sheet_records: list[JsonDict], rows: list[JsonDict]) -> JsonDict:
    errors = []
    counts = Counter()
    by_sheet = {record["sheet"]: record for record in sheet_records}
    rows_by_sheet: dict[str, list[JsonDict]] = {}
    for row in rows:
        rows_by_sheet.setdefault(row["sheet"], []).append(row)
    for sheet in workbook.worksheets:
        record = by_sheet[sheet.title]
        if record["mapping_status"] == "mapping_required":
            continue
        nonempty = [cell for line in sheet for cell in line if not isinstance(cell, MergedCell) and cell.value is not None]
        counts["raw_cells_verified"] += len(nonempty)
        for field, mapped in record["fields"].items():
            if mapped["header_cell_address"] is None:
                continue
            if sheet[mapped["header_cell_address"]].value != mapped["original_header_text"]:
                errors.append(f"Header mapping mismatch: {sheet.title} {field}")
            else:
                counts["header_fields_verified"] += 1
        for row in rows_by_sheet.get(sheet.title, []):
            for field in FIELDS:
                provenance = row["field_provenance"][field]
                if provenance["resolution"] == "unmapped":
                    continue
                raw_cell = sheet[provenance["source_cell"]]
                raw_value = _serializable(raw_cell.value if not isinstance(raw_cell, MergedCell) else None)[0]
                if (raw_value != provenance["original_value"]
                        and (field != "project" or provenance["resolution"] not in ("hierarchical_project_group", "unresolved_group"))):
                    errors.append(f"Field provenance mismatch: {sheet.title}!{provenance['source_cell']} {field}")
                if field in ("actual", "next_plan", "issue_risk", "pic") and row[field] != raw_value:
                    errors.append(f"Content changed: {sheet.title}!{provenance['source_cell']} {field}")
                if field == "project" and provenance["resolution"] == "hierarchical_project_group":
                    if row[field] != _serializable(sheet[provenance["context_cell"]].value)[0]:
                        errors.append(f"Project label mismatch: {sheet.title}!{provenance['context_cell']}")
                    if sheet[provenance["context_cell"]].row > row["source_row"]:
                        errors.append(f"Future project label: {sheet.title} row {row['source_row']}")
                if not row["source_refs"].get(field):
                    errors.append(f"Missing source ref: {sheet.title} row {row['source_row']} {field}")
                counts["normalized_fields_verified"] += 1
    mapping_required = any(record["mapping_status"] == "mapping_required" for record in sheet_records)
    status = "mapping_required" if mapping_required else ("PASS" if not errors else "FAIL")
    return {"status": status, "counts": dict(counts), "errors": errors}


def inspect_workbook(
    path: Path, *, mapping_overrides: dict[str, JsonDict] | None = None
) -> WorkbookImportPreview:
    path = Path(path)
    workbook_sha256 = hashlib.sha256(path.read_bytes()).hexdigest()
    workbook = load_workbook(path, data_only=False, read_only=False)
    sheets, rows, excluded_rows = [], [], {}
    overrides = mapping_overrides or {}
    applied_overrides = set()
    try:
        for index, sheet in enumerate(workbook.worksheets):
            override = overrides.get(sheet.title)
            if override is not None:
                config, fields = _manual_mapping(sheet, override)
                period = _period_of(sheet, config)
                sheet_rows, exclusions = _manual_flat_rows(sheet, config)
                state = ("draft_or_incomplete" if sheet_rows and not any(
                    _meaningful(item[field]) for item in sheet_rows for field in ("actual", "next_plan", "issue_risk")
                ) else "not_assessed")
                for row in sheet_rows:
                    row["reporting_period"] = period["label"]
                    row["reporting_year"] = period["year"]
                rows.extend(sheet_rows)
                excluded_rows[sheet.title] = exclusions
                sheets.append({
                    "sheet_index": index,
                    "sheet": sheet.title,
                    "template": config["template_id"],
                    "mapping_status": "confirmed",
                    "mapping_version": config["version"],
                    "header_rows": [config["header_row"]],
                    "column_header_row": config["header_row"],
                    "first_data_row": config["first_data_row"],
                    "sheet_state": state,
                    "reporting_period": period,
                    "fields": fields,
                    "ambiguities": ["Explicit flat mapping; blank project/role labels are not inferred"],
                    "normalized_row_count": len(sheet_rows),
                    "user_mapping": override,
                })
                applied_overrides.add(sheet.title)
                continue
            config = _get_template(sheet)
            if config is None:
                sheets.append({
                    "sheet_index": index, "sheet": sheet.title, "template": None,
                    "mapping_status": "mapping_required", "sheet_state": "mapping_required",
                    "reporting_period": None, "fields": {},
                    "ambiguities": ["Unknown worksheet layout; mapping is required"],
                    "normalized_row_count": 0,
                })
                excluded_rows[sheet.title] = []
                continue

            merges = _merge_index(sheet)
            period = _period_of(sheet, config)
            fields = _header_mapping(sheet, config, merges)
            sheet_rows, exclusions = _sheet_rows(sheet, config, merges)
            first_data_row = sheet_rows[0]["source_row"] if sheet_rows else None
            state = ("draft_or_incomplete" if sheet_rows and not any(
                _meaningful(item[field]) for item in sheet_rows for field in ("actual", "next_plan", "issue_risk")
            ) else "not_assessed")
            for mapping in fields.values():
                mapping["first_data_row"] = first_data_row
            for row in sheet_rows:
                row["reporting_period"] = period["label"]
                row["reporting_year"] = period["year"]
            rows.extend(sheet_rows)
            excluded_rows[sheet.title] = exclusions
            ambiguities = ["Project and role share a physical header; use the recorded hierarchy"]
            if config["name"] == "revised":
                ambiguities.append("Column D can contain a more-specific role without an independent header")
            ambiguities.extend(period["conflicts"])
            sheets.append({
                "sheet_index": index, "sheet": sheet.title, "template": config["name"],
                "mapping_status": "recognized", "header_rows": [1, 2, 4] if config["name"] == "early" else [2, 3, 5],
                "column_header_row": config["header_row"], "first_data_row": first_data_row,
                "sheet_state": state, "reporting_period": period, "fields": fields,
                "ambiguities": ambiguities, "normalized_row_count": len(sheet_rows),
            })
        unknown_overrides = set(overrides) - applied_overrides
        if unknown_overrides:
            raise ValueError(f"Manual mapping references unknown sheets: {sorted(unknown_overrides)}")
        validation = _validate_extraction(workbook, sheets, rows)
    finally:
        workbook.close()
    return WorkbookImportPreview(
        workbook_sha256=workbook_sha256,
        workbook_filename=path.name,
        sheets=sheets,
        rows=rows,
        excluded_rows=excluded_rows,
        validation=validation,
    )


def normalized_payload(
    preview: WorkbookImportPreview, *, workbook_filename: str | None = None
) -> JsonDict:
    """Build the stable normalized-input envelope used by historical replay hashes."""
    return {
        "schema_version": "1.0",
        "workbook_filename": preview.workbook_filename if workbook_filename is None else workbook_filename,
        "source_sha256": preview.workbook_sha256,
        "rows": [{key: value for key, value in row.items() if key != "reporting_year"} for row in preview.rows],
        "excluded_rows": preview.excluded_rows,
        "validation": preview.validation,
    }
